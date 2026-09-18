from __future__ import annotations

import base64
import json
import tempfile
import unittest
import uuid
import time
import httpx2 as httpx
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import sessionmaker
from docx import Document

from review_writer_api.app import create_app
from review_writer_api.config import ApiSettings
from review_writer_api.database import Base, Project, User
from review_writer_api.domain_services.planning import (
    MATRIX_FACT_ENRICHMENT_CONTRACT_VERSION,
    MATRIX_FACT_PROMPT_VERSION,
    ROUTING_REQUIRED_LABEL,
    _matrix_classification_axes,
    _required_topic_partitions_from_outline,
    _topic_outline_intent,
    _topic_partition_for_row,
    _topic_partition_routes,
    _usable_fact_candidate,
)
from review_writer_api.domain_services.library_index import EvidenceHit
from review_writer_core.stages.planning.topic import TOPIC_PARTITION_BOUNDARY_LABEL, _topic_partition_for_text
from review_writer_api.security import Principal, Role
from review_writer_api.errors import WorkflowConflict, WorkflowValidationError
from review_writer_api.workflow_models import LibraryPaper
from review_writer_core.scientific_facts import FACT_PROMPT_VERSION
from review_writer_api.tests.planning_fixtures import seed_verified_matrix, offline_argument_planner


TEST_KEY = base64.urlsafe_b64encode(bytes(range(32))).decode("ascii").rstrip("=")


class PlanningV1Tests(unittest.TestCase):
    def test_missing_topic_recommendation_is_a_persisted_candidate_not_an_outline_write(self):
        from review_writer_core.stages.planning.topic_recommendation import validate_recommendation
        calls = []
        def build(context, payload):
            calls.append(payload)
            return validate_recommendation({"organization": "Scientific approaches", "sections": [
                {"title": "Approaches and evidence", "role": "body", "question": "How do the approaches differ?",
                 "paper_ids": [p["paper_id"] for p in payload["papers"]], "rationale": "Compare the selected studies."}
            ]}, payload["papers"])
        service = self.app.state.planning_service
        with patch.dict(self.app.state.job_service._handlers, {"planning.topic-outline": build}), \
             patch("review_writer_api.domain_services.planning._topic_outline_intent", return_value={"available": False}), \
             TestClient(self.app) as client:
            base = f"/api/v1/projects/{self.project_id}/planning"
            before = client.get(base).json()
            self.assertEqual("not_started", before["topic_recommendation"]["status"])
            self.assertEqual([], calls)  # GET never initiates paid work.
            queued = client.post(base + "/outline/topic/jobs", headers=self.headers())
            self.assertEqual(202, queued.status_code, queued.text)
            job = queued.json()
            for _ in range(100):
                job = client.get(f"/api/v1/jobs/{job['id']}").json()
                if job["status"] not in {"queued", "running"}: break
                time.sleep(.03)
            self.assertEqual("succeeded", job["status"], job)
            current = client.get(base).json()
            self.assertEqual(before["outline_selection"], current["outline_selection"])
            self.assertTrue(any(c["source"] == "topic" for c in current["outline_candidates"]))
            self.assertEqual(job["id"], client.post(base + "/outline/topic/jobs", headers=self.headers()).json()["id"])
            self.assertEqual(1, len(calls))
            snapshot = next(c["outline_md"] for c in current["outline_candidates"] if c["source"] == "topic")
            rejected = client.put(base + "/outline", headers=self.headers(), json={"revision": current["matrix_revision"], "outline_style": "topic-guided", "candidate_outline_md": snapshot + "changed"})
            self.assertEqual(409, rejected.status_code)
            with patch.object(service, "_outline_sources", side_effect=AssertionError("Applying the cached outline must not reread sources")):
                chosen = client.put(base + "/outline", headers=self.headers(), json={"revision": current["matrix_revision"], "outline_style": "topic-guided", "candidate_outline_md": snapshot})
            self.assertEqual(200, chosen.status_code, chosen.text)
            self.assertEqual(snapshot.strip(), chosen.json()["selected_outline_md"].strip())
            self.assertIn("Approaches and evidence", client.get(base).json()["selected_outline_md"])
            self.current = self.second
            self.assertEqual(404, client.post(base + "/outline/topic/jobs", headers=self.headers()).status_code)

    def test_candidate_source_failure_is_local_and_retains_valid_previous_fact(self):
        from copy import deepcopy
        from review_writer_core.scientific_facts import fact_is_usable
        service = self.app.state.planning_service
        seed_verified_matrix(service, self.first, self.project_id)
        original, _ = service._matrix(self.first, self.project_id)
        candidate = deepcopy(original)
        row = candidate["rows"][0]
        old = deepcopy(row["scientific_facts"][0])
        row["scientific_facts"][0]["value"] = "Unsupported changed conclusion 99%."
        added = deepcopy(old)
        added.update(fact_id="NEW-UNSUPPORTED", value="Unverified new result 98%.")
        row["scientific_facts"].append(added)
        issues = service._validate_candidate_matrix(self.first, original, candidate)
        self.assertEqual(["retained_previous", "withheld"], [i["action"] for i in issues])
        self.assertEqual(old, row["scientific_facts"][0])
        self.assertFalse(fact_is_usable(row["scientific_facts"][1]))
        self.assertEqual(original["rows"][1]["scientific_facts"], candidate["rows"][1]["scientific_facts"])
        self.assertIn("verification_outdated", issues[0]["reasons"])
        self.assertEqual([], service._validate_candidate_matrix(self.first, original, candidate))
        missing = deepcopy(candidate)
        missing["rows"].pop()
        with self.assertRaises(WorkflowValidationError):
            service._validate_candidate_matrix(self.first, original, missing)

    def test_withheld_fact_candidate_can_be_planned_published_and_confirmed(self):
        from copy import deepcopy
        from review_writer_core.scientific_facts import fact_is_usable
        service = self.app.state.planning_service
        seed_verified_matrix(service, self.first, self.project_id)
        with TestClient(self.app) as client:
            self.choose_outline(client, "reaction")
            prepared = service.blueprint_job_payload(self.first, self.project_id, revision=0)
            added = deepcopy(prepared["matrix_snapshot"]["rows"][0]["scientific_facts"][0])
            added.update(fact_id="NEW-UNSUPPORTED", value="Unverified result 99%.")
            prepared["matrix_snapshot"]["rows"][0]["scientific_facts"].append(added)
            prepared["planning_matrix_snapshot"] = deepcopy(prepared["matrix_snapshot"])
            service.reconcile_blueprint_facts(self.first, self.project_id, prepared)
            self.assertFalse(fact_is_usable(prepared["planning_matrix_snapshot"]["rows"][0]["scientific_facts"][-1]))
            built = offline_argument_planner(None, prepared)
            for section in built["section_blueprint"]["sections"]:
                self.assertNotIn("NEW-UNSUPPORTED", (section.get("planning_fact_context") or {}).get("fact_ids", []))
            result = service.publish_blueprint_candidate(self.first, self.project_id, built)
            self.assertEqual("withheld", result["section_blueprint"]["fact_source_issues"][0]["action"])
            confirmed = service.confirm_blueprint(self.first, self.project_id, revision=0, artifact_id=result["blueprint_artifact_id"])
            self.assertEqual("approved", confirmed["status"])

    def test_fact_supplement_uses_linked_si_without_changing_citation_identity(self):
        service = self.app.state.planning_service
        with self.sessions.begin() as session:
            linked = session.scalar(select(LibraryPaper).where(LibraryPaper.user_id == uuid.UUID(self.first.user_id), LibraryPaper.paper_id == "P002"))
            linked.metadata_json = {**linked.metadata_json, "document_type": "supporting_information", "parent_paper_id": "P001"}
        current = service.get(self.first, self.project_id)
        paper = {"paper_id": "P001", "source_lineages": {"P001": "main", "P002": "si"}, "evidence_candidates": []}
        payload = {"source_matrix_artifact_id": current["matrix_artifact_id"], "papers": [paper]}
        hit = EvidenceHit(paper_id="P002", chunk_id="table", content="Entry 1 gave 91% yield.",
            page_start=3, page_end=3, section_path=("Experimental",), content_type="table", asset_refs=(),
            score=1, match_reason="test", is_neighbor=False, index_id="idx", source_lineage_hash="si")
        previous = service.library_index
        index = SimpleNamespace(enabled=True, summaries=lambda *_: {"P001": {"source_lineage_hash": "main"}, "P002": {"source_lineage_hash": "si"}},
            retrieve=lambda *a, **k: [hit])
        service.library_index = index
        try:
            found = service.retrieve_matrix_fact_evidence(self.first, self.project_id, payload,
                {"paper_id": "P001", "questions": [{"field_id": "quantitative_results", "query": "yield"}]})
            self.assertEqual("P001", found[0]["paper_id"])
            self.assertEqual("P002", found[0]["source_file_id"])
            self.assertEqual("supporting_information", found[0]["source_type"])
            with self.sessions.begin() as session:
                linked = session.scalar(select(LibraryPaper).where(LibraryPaper.user_id == uuid.UUID(self.first.user_id), LibraryPaper.paper_id == "P002"))
                linked.metadata_json = {**linked.metadata_json, "parent_paper_id": "P003"}
            with self.assertRaises(WorkflowConflict):
                service._validate_fact_sources(self.first, [paper])
        finally:
            service.library_index = previous

    def test_agent_retrieval_registers_only_current_owned_paper_sources(self):
        service = self.app.state.planning_service
        current = service.get(self.first, self.project_id)
        payload = {"source_matrix_artifact_id": current["matrix_artifact_id"], "papers": [
            {"paper_id": "P001", "index_summary": {"source_lineage_hash": "lineage"}, "evidence_candidates": []}
        ]}
        def hit(paper, lineage):
            return EvidenceHit(paper_id=paper, chunk_id="results", content="The result was 91% yield.",
                page_start=2, page_end=2, section_path=("Results",), content_type="text", asset_refs=(),
                score=1, match_reason="test", is_neighbor=False, index_id="idx", source_lineage_hash=lineage)
        index = SimpleNamespace(enabled=True, summaries=lambda _principal, _ids: {"P001": {"source_lineage_hash": "lineage"}},
            retrieve=lambda *a, **k: [hit("P001", "lineage"), hit("P002", "lineage"), hit("P001", "old")])
        previous = service.library_index
        service.library_index = index
        try:
            request = {"paper_id": "P001", "questions": [{"field_id": "quantitative_results", "query": "reported yield"}]}
            found = service.retrieve_matrix_fact_evidence(self.first, self.project_id, payload, request)
            self.assertEqual(1, len(found))
            self.assertEqual(found, payload["papers"][0]["evidence_candidates"])
            with self.assertRaises(WorkflowValidationError):
                service.retrieve_matrix_fact_evidence(self.first, self.project_id, payload, {**request, "paper_id": "P002"})
            index.summaries = lambda _principal, _ids: {"P001": {"source_lineage_hash": "changed"}}
            with self.assertRaises(WorkflowConflict):
                service.retrieve_matrix_fact_evidence(self.first, self.project_id, payload, request)
        finally:
            service.library_index = previous

    def test_fact_agent_recovers_long_question_with_one_semantic_pass(self):
        service = self.app.state.planning_service
        current = service.get(self.first, self.project_id)
        payload = {"source_matrix_artifact_id": current["matrix_artifact_id"], "papers": [
            {"paper_id": "P001", "index_summary": {"source_lineage_hash": "lineage"}, "evidence_candidates": []}]}
        calls = []
        def retrieve(principal, query, **kwargs):
            calls.append((query, kwargs))
            if len(calls) < 3:
                return []
            return [EvidenceHit(paper_id="P001", chunk_id="table", content="External cohort accuracy was 84%.",
                page_start=3, page_end=3, section_path=("Validation",), content_type="table", asset_refs=(),
                score=1, match_reason="test", is_neighbor=False, index_id="idx", source_lineage_hash="lineage")]
        old = service.library_index
        service.library_index = SimpleNamespace(enabled=True, summaries=lambda *_: {"P001": {"source_lineage_hash": "lineage"}}, retrieve=retrieve)
        try:
            question = "Please identify all the quantitative results for the external cohort and surrounding discussion."
            found = service.retrieve_matrix_fact_evidence(self.first, self.project_id, payload, {"paper_id": "P001", "questions": [
                {"field_id": "quantitative_results", "query": question, "target_terms": ["external cohort", "accuracy"]}]})
            self.assertEqual(1, len(found))
            self.assertEqual([True, False, False], [kwargs["use_semantic"] for _, kwargs in calls])
            self.assertTrue(all(kwargs["allowed_papers"] == ["P001"] for _, kwargs in calls))
            self.assertTrue(all(query != question for query, _ in calls))
            self.assertEqual(question, calls[0][1]["semantic_query"])
        finally:
            service.library_index = old

    def test_fact_source_recovery_uses_only_registered_pages_and_retains_provenance(self):
        service = self.app.state.planning_service
        current = service.get(self.first, self.project_id)
        payload = {"source_matrix_artifact_id": current["matrix_artifact_id"], "papers": [
            {"paper_id": "P001", "index_summary": {"source_lineage_hash": "lineage"}, "evidence_candidates": [
                {"evidence_key": "registered", "paper_id": "P001", "source_file_id": "P001", "page_start": 3,
                 "content": "The temperature was damaged in extraction.", "question_ids": ["method_conditions"]}]}]}
        recovered = EvidenceHit(paper_id="P001", chunk_id="pdf-text:current:p3", content="The annealing temperature was 245 K.",
            page_start=3, page_end=3, section_path=("Registered PDF text layer",), content_type="pdf_text", asset_refs=(),
            score=0, match_reason="source_text_recovery", is_neighbor=False, index_id="", source_lineage_hash="lineage")
        calls = []
        def recover(principal, paper_id, pages, **kwargs):
            calls.append((paper_id, pages, kwargs["expected_lineage"]))
            return [recovered]
        old = service.library_index
        service.library_index = SimpleNamespace(enabled=True, summaries=lambda *_: {"P001": {"source_lineage_hash": "lineage"}},
            recover_pdf_pages=recover, retrieve=lambda *a, **k: self.fail("Do not query the index again after recovering new PDF text"))
        try:
            found = service.retrieve_matrix_fact_evidence(self.first, self.project_id, payload, {"paper_id": "P001", "questions": [
                {"field_id": "method_conditions", "query": "Exact experimental temperature", "source_recovery": True,
                 "evidence_keys": ["registered", "unregistered-foreign-page"], "page_start": 999}]})
            self.assertEqual([("P001", [3], "lineage")], calls)
            self.assertEqual(1, len(found))
            self.assertEqual("P001", found[0]["source_file_id"])
            self.assertEqual("lineage", found[0]["source_lineage_hash"])
            self.assertEqual("pdf_text", found[0]["content_type"])
            self.assertEqual(2, len(payload["papers"][0]["evidence_candidates"]))
        finally:
            service.library_index = old

    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        root = Path(self.temporary.name)
        database_url = f"sqlite+pysqlite:///{(root / 'planning.sqlite3').as_posix()}"
        self.engine = create_engine(database_url, connect_args={"check_same_thread": False})

        @event.listens_for(self.engine, "connect")
        def enable_foreign_keys(connection, _record):
            cursor = connection.cursor()
            cursor.execute("PRAGMA foreign_keys=ON")
            cursor.close()

        Base.metadata.create_all(self.engine)
        self.sessions = sessionmaker(bind=self.engine, expire_on_commit=False)
        with self.sessions.begin() as session:
            first = User(email="first@example.com", display_name="First", password_hash="hash")
            second = User(email="second@example.com", display_name="Second", password_hash="hash")
            session.add_all([first, second])
            session.flush()
            project = Project(
                user_id=first.id,
                slug="planning",
                topic="Copper allenation",
                taxonomy_profile="chemistry_general",
            )
            hidden = Project(user_id=second.id, slug="hidden", topic="Hidden")
            session.add_all([project, hidden])
            session.flush()
            papers = []
            substrates = (
                "aromatic substrates",
                "small-molecule substrates",
                "biomolecular substrates",
            )
            catalysts = (
                "transition-metal catalysis",
                "organocatalysis",
                "photochemical methods",
            )
            reactions = (
                "cross-coupling",
                "addition reactions",
                "cyclization and annulation",
            )
            for index in range(1, 36):
                paper_id = f"P{index:03d}"
                structured_tags = {
                    "substrate": substrates[(index - 1) % len(substrates)],
                    "catalyst_or_method": catalysts[(index - 1) % len(catalysts)],
                    "reaction_type": reactions[(index - 1) % len(reactions)],
                }
                papers.append(
                    LibraryPaper(
                        user_id=first.id,
                        paper_id=paper_id,
                        content_sha256=f"{index:064x}",
                        original_filename=f"{paper_id}.pdf",
                        title=f"Paper {index}",
                        authors_json=[f"Author {index}"],
                        keywords_json=["allenation", "copper"],
                        tags_json=structured_tags,
                        metadata_json={
                            "paper_id": paper_id,
                            "title": {"value": f"Paper {index}"},
                            "authors": {"value": [f"Author {index}"]},
                            "keywords": {"value": ["allenation", "copper"]},
                            "abstract": {"value": f"Evidence for {paper_id}."},
                            "year": {"value": 2024},
                            "structured_tags": {
                                "value": structured_tags,
                                "human_checked": True,
                            },
                        },
                        pdf_relative_path=f"review-library/uploads/{paper_id}.pdf",
                        markdown_relative_path=f"review-library/markdown/{paper_id}.md",
                    )
                )
            session.add_all(papers)
            self.project_id = str(project.id)
            self.hidden_project_id = str(hidden.id)
            self.first = Principal(str(first.id), frozenset({Role.USER}), first.email)
            self.second = Principal(str(second.id), frozenset({Role.USER}), second.email)
        self.current = self.first
        settings = ApiSettings(
            review_root=root,
            deployment_mode="hosted",
            database_url=database_url,
            public_origin="http://testserver",
            credential_encryption_key=TEST_KEY,
            hosted_workspace_root=root / "users",
        )
        self.app = create_app(
            settings,
            principal_provider=lambda: self.current,
            session_factory_override=self.sessions,
            native_workflow_overrides={"planning.blueprint": offline_argument_planner},
        )
        self._seed_discovery(range(1, 36))

    def tearDown(self) -> None:
        Base.metadata.drop_all(self.engine)
        self.engine.dispose()
        self.temporary.cleanup()

    def generate_candidate(self, client, **kwargs):
        """Exercise the async endpoint, then expose its completed candidate to assertions."""
        response = client.post(f"/api/v1/projects/{self.project_id}/planning/blueprint", **kwargs)
        if response.status_code != 202:
            return response
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            job = client.get(f"/api/v1/jobs/{response.json()['id']}").json()
            if job["status"] in {"succeeded", "failed", "cancelled"}:
                self.assertEqual("succeeded", job["status"], job)
                return httpx.Response(200, json=job["result"])
            time.sleep(0.02)
        self.fail("Planning job did not finish.")

    @staticmethod
    def headers() -> dict[str, str]:
        return {"Origin": "http://testserver"}

    def _review(self, selected: set[int]) -> dict:
        return {
            "project_id": self.project_id,
            "topic": "Copper allenation",
            "selection_mode": "explicit",
            "results": [
                {
                    "keyword": "allenation",
                    "keep": True,
                    "local_results": [
                        {
                            "paper_id": f"P{index:03d}",
                            "title": f"Paper {index}",
                            "score": 100 - index,
                            "role": "core_candidate",
                            "selected_for_matrix": index in selected,
                        }
                        for index in range(1, 36)
                    ],
                    "web_results": [],
                }
            ],
        }

    def _seed_discovery(self, selected) -> None:
        service = self.app.state.discovery_service
        repository = self.app.state.workflow_repository
        artifact, run = service._write_json_artifact(
            self.first,
            self.project_id,
            stage_id="discovery",
            logical_name="discovery/review.json",
            payload=self._review(set(selected)),
            make_current=False,
        )
        repository.save_discovery_atomically(
            self.first.user_id,
            self.project_id,
            artifact_id=artifact.id,
            run_id=run.id,
            expected_revision=0,
            status="review",
        )
        with TestClient(self.app) as client:
            response = client.post(
                f"/api/v1/projects/{self.project_id}/discovery/confirm",
                json={"revision": 1},
                headers=self.headers(),
            )
        self.assertEqual(200, response.status_code, response.text)

    def planning(self, client: TestClient) -> dict:
        response = client.get(f"/api/v1/projects/{self.project_id}/planning")
        self.assertEqual(200, response.status_code, response.text)
        return response.json()

    def choose_outline(self, client: TestClient, style: str = "substrate") -> dict:
        current = self.planning(client)
        response = client.put(
            f"/api/v1/projects/{self.project_id}/planning/outline",
            json={"revision": current["matrix_revision"], "outline_style": style},
            headers=self.headers(),
        )
        self.assertEqual(200, response.status_code, response.text)
        return response.json()

    def test_real_matrix_publication_handoff_reuses_extraction_cache(self):
        service, repo = self.app.state.planning_service, self.app.state.workflow_repository
        with TestClient(self.app) as client:
            self.choose_outline(client, "reaction")
            source = service.matrix_enrichment_payload(self.first, self.project_id)
            job = repo.create_or_get_job(self.first.user_id, self.project_id, "project", "matrix.enrich",
                "handoff", {"source_matrix_artifact_id": source["source_matrix_artifact_id"]},
                operation_key="matrix-enrichment")
            prepared = service.blueprint_job_payload(self.first, self.project_id, revision=0)
            self.assertEqual(job.id, prepared["await_matrix_job_id"])
            claimed = repo.claim_job(job.id)
            result = service.publish_matrix_enrichment(self.first, self.project_id, source,
                {"papers": [{"paper_id": paper["paper_id"], "status": "complete",
                    "fact_extraction_profile": {"stop_reason": "checks_completed"},
                    "facts": [], "failed_fields": []} for paper in source["papers"]]})
            repo.mark_job_succeeded(job.id, result, lease_token=claimed.lease_token,
                lease_generation=claimed.lease_generation)
            refreshed = service.resume_blueprint_after_matrix(self.first, self.project_id, prepared)
            self.assertEqual(result["matrix_artifact_id"], refreshed["section_blueprint"]["source_matrix_artifact_id"])
            self.assertEqual(0, service.matrix_enrichment_payload(self.first, self.project_id)["pending_paper_count"])
            service.validate_prepared_blueprint(self.first, self.project_id, refreshed)

    def test_fact_contract_version_invalidates_old_cache_and_force_can_reextract(self) -> None:
        service = self.app.state.planning_service
        source = service.matrix_enrichment_payload(self.first, self.project_id)
        self.assertEqual(
            MATRIX_FACT_ENRICHMENT_CONTRACT_VERSION,
            source["fact_enrichment_contract_version"],
        )
        service.publish_matrix_enrichment(
            self.first,
            self.project_id,
            source,
            {
                "papers": [
                    {
                        "paper_id": paper["paper_id"],
                        "status": "complete",
                        "fact_extraction_profile": {"stop_reason": "checks_completed"},
                        "facts": [],
                        "failed_fields": [],
                    }
                    for paper in source["papers"]
                ]
            },
        )

        current = service.matrix_enrichment_payload(self.first, self.project_id)
        forced = service.matrix_enrichment_payload(
            self.first,
            self.project_id,
            force=True,
        )

        self.assertEqual(0, current["pending_paper_count"])
        self.assertEqual(source["paper_count"], forced["pending_paper_count"])
        self.assertTrue(forced["force_refresh"])

        with self.sessions.begin() as session:
            project = session.get(Project, uuid.UUID(self.project_id))
            project.model_tier = "luna"
        changed_model = service.matrix_enrichment_payload(
            self.first,
            self.project_id,
        )
        self.assertEqual(
            source["paper_count"], changed_model["pending_paper_count"]
        )
        self.assertEqual("gpt-5.6-luna", changed_model["actual_model_id"])

    def test_matrix_extraction_axes_ignore_runtime_coverage_fields(self) -> None:
        matrix = {
            "classification_axes": [
                {
                    "axis_id": "axis_01",
                    "label": "Method family",
                    "axis_role": "primary_organization",
                    "role_status": "evidence_confirmed",
                    "evidence_coverage": {"paper_count": 12, "partition_count": 3},
                    "partitions": [
                        {"partition_id": "method_a", "label": "Method A"}
                    ],
                }
            ]
        }

        axes = _matrix_classification_axes(matrix, [])

        self.assertNotIn("role_status", axes[0])
        self.assertNotIn("evidence_coverage", axes[0])
        self.assertIn("role_status", matrix["classification_axes"][0])

    def test_matrix_fact_retrieval_recovers_question_hits_inside_admitted_paper(self) -> None:
        class RecoveryIndex:
            enabled = True

            @staticmethod
            def summaries(_principal, paper_ids):
                return {
                    paper_id: {
                        "fulltext": "ready",
                        "chunker_version": "test-v1",
                        "source_lineage_hash": f"lineage-{paper_id}",
                    }
                    for paper_id in paper_ids
                    if paper_id == "P001"
                }

            @staticmethod
            def retrieve(
                _principal,
                _query,
                *,
                allowed_papers,
                term_groups=None,
                **_kwargs,
            ):
                groups = list(term_groups or [])
                if len(groups) != 1 or "substrate" not in groups[0]:
                    return []
                paper_id = allowed_papers[0]
                return [
                    EvidenceHit(
                        paper_id=paper_id,
                        chunk_id="results-001",
                        content="The substrate scope included aromatic examples.",
                        page_start=3,
                        page_end=3,
                        section_path=("Results",),
                        content_type="text",
                        asset_refs=(),
                        score=1.0,
                        match_reason="test",
                        is_neighbor=False,
                        index_id="index-001",
                        source_lineage_hash=f"lineage-{paper_id}",
                    )
                ]

        service = self.app.state.planning_service
        previous = service.library_index
        service.library_index = RecoveryIndex()
        try:
            payload = service.matrix_enrichment_payload(
                self.first, self.project_id, force=True
            )
        finally:
            service.library_index = previous

        paper = next(
            item for item in payload["papers"] if item["paper_id"] == "P001"
        )
        recovered = next(
            item
            for item in paper["evidence_candidates"]
            if item.get("chunk_id") == "results-001"
        )
        self.assertIn("object_input", recovered["question_ids"])
        self.assertIn(
            "admitted_paper_question_recovery", recovered["retrieval_passes"]
        )
        self.assertGreater(
            paper["retrieval_summary"]["relaxed_question_hit_count"], 0
        )
        self.assertEqual(payload["required_fact_roles"], paper["required_fact_roles"])
        self.assertEqual([], paper["required_fact_roles"])
        self.assertEqual(FACT_PROMPT_VERSION, MATRIX_FACT_PROMPT_VERSION)

    def test_path_only_mineru_image_is_not_a_fact_candidate(self) -> None:
        self.assertFalse(
            _usable_fact_candidate("image", "images/8a42f1ac87f2.png")
        )
        self.assertFalse(
            _usable_fact_candidate("text", r"figures\scheme-1.jpeg")
        )
        self.assertTrue(
            _usable_fact_candidate(
                "image",
                "Scheme 1 shows the reported catalytic conversion.",
            )
        )

    def test_prior_matrix_facts_remain_available_as_provider_failure_fallback(self) -> None:
        service = self.app.state.planning_service
        artifact, _run = self.app.state.discovery_service._write_json_artifact(
            self.first,
            self.project_id,
            stage_id="matrix",
            logical_name="matrix/fallback-test.json",
            payload={
                "rows": [
                    {
                        "paper_id": "P001",
                        "scientific_facts": [
                            {
                                "fact_id": "MF-PREVIOUS",
                                "field_id": "object_input",
                                "value": "Terminal alkynes and aldehydes.",
                                "evidence_refs": [
                                    {"evidence_key": "sha256:previous"}
                                ],
                            },
                            {
                                "fact_id": "MF-CLASSIFICATION",
                                "field_id": "topic_partition",
                                "value": "A runtime classification.",
                                "evidence_refs": [
                                    {"evidence_key": "sha256:previous"}
                                ],
                            },
                        ],
                    }
                ]
            },
            make_current=False,
        )

        fallback = service._prior_matrix_facts(
            self.first,
            {
                "fact_enrichment_summary": {
                    "source_matrix_artifact_id": artifact.id
                }
            },
        )

        self.assertEqual(["MF-PREVIOUS"], [row["fact_id"] for row in fallback["P001"]])

    def test_current_fact_refresh_is_an_idempotent_success_not_an_error(self) -> None:
        service = self.app.state.planning_service
        with patch.object(
            service,
            "matrix_enrichment_payload",
            return_value={
                "project_id": self.project_id,
                "source_matrix_artifact_id": service._matrix(self.first, self.project_id)[1].id,
                "pending_paper_count": 0,
                "fulltext_candidate_paper_count": 0,
            },
        ), TestClient(self.app) as client:
            response = client.post(
                f"/api/v1/projects/{self.project_id}/planning/matrix/enrichment/jobs",
                headers={**self.headers(), "Idempotency-Key": "current-facts"},
            )

        self.assertEqual(202, response.status_code, response.text)
        self.assertEqual("queued", response.json()["status"])
        # The cache check is deferred to the worker, not performed by POST.
        job = self.app.state.workflow_repository.get_job(self.first.user_id, response.json()["id"])
        self.assertTrue(job.payload["prepare_on_start"])

    def test_fact_retry_route_passes_only_requested_papers(self) -> None:
        service = self.app.state.planning_service
        with patch.object(service, "matrix_enrichment_payload", return_value={"pending_paper_count": 0}) as prepare, TestClient(self.app) as client:
            response = client.post(
                f"/api/v1/projects/{self.project_id}/planning/matrix/enrichment/jobs?paper_ids=P001&force=true",
                headers=self.headers(),
            )
        self.assertEqual(202, response.status_code, response.text)
        prepare.assert_not_called()
        job = self.app.state.workflow_repository.get_job(self.first.user_id, response.json()["id"])
        self.assertEqual(["P001"], job.payload["selected_paper_ids"])
        self.assertTrue(job.payload["force_refresh"])

    def test_failed_fact_retry_keeps_valid_same_source_result(self) -> None:
        service = self.app.state.planning_service
        matrix, artifact = service._matrix(self.first, self.project_id)
        row = matrix["rows"][0]
        fact = {"fact_id": "previous", "field_id": "object_input", "value": "Source result",
                "evidence_refs": [{"evidence_key": "original"}]}
        row["scientific_facts"] = [fact]
        row["fact_enrichment"] = {"status": "complete", "source_fingerprint": "same-source"}
        payload = {"papers": [{"paper_id": row["paper_id"], "source_fingerprint": "same-source"}]}
        built = {"papers": [{"paper_id": row["paper_id"], "status": "failed", "facts": [], "error": "provider timed out"}]}
        with patch.object(service, "validate_matrix_enrichment_inputs", return_value=(matrix, artifact)):
            result = service.publish_matrix_enrichment(self.first, self.project_id, payload, built, candidate_only=True)
        restored = result["matrix_snapshot"]["rows"][0]
        self.assertEqual([fact], restored["scientific_facts"])
        self.assertEqual("complete", restored["fact_enrichment"]["status"])
        self.assertEqual("failed", restored["fact_enrichment"]["last_attempt"]["status"])

    def test_all_fact_failures_allow_provisional_planning_without_limited_mode(self) -> None:
        service = self.app.state.planning_service
        current = service.get(self.first, self.project_id)
        rows = current["literature_matrix"]["rows"]
        source_payload = {
            "source_matrix_artifact_id": current["matrix_artifact_id"],
            "expected_matrix_revision": current["matrix_revision"],
            "papers": [
                {
                    "paper_id": row["paper_id"],
                    "source_fingerprint": f"failed-{row['paper_id']}",
                    "index_summary": {},
                    "evidence_candidates": [],
                }
                for row in rows
            ],
        }
        built = {
            "papers": [
                {
                    "paper_id": row["paper_id"],
                    "status": "failed",
                    "facts": [],
                    "failed_fields": ["all"],
                    "error": "provider unavailable",
                }
                for row in rows
            ]
        }
        service.publish_matrix_enrichment(
            self.first, self.project_id, source_payload, built
        )

        with TestClient(self.app) as client:
            blocked = self.planning(client)
            self.assertFalse(blocked["matrix_enrichment"]["planning_blocked"])
            self.choose_outline(client)
            blocked = self.planning(client)
            response = self.generate_candidate(client,
                json={"revision": blocked["blueprint_revision"]},
                headers=self.headers(),
            )
            self.assertEqual(200, response.status_code, response.text)
            candidate = response.json()
            self.assertEqual("completed", candidate["section_blueprint"]["academic_planning"]["status"])
            confirmed = client.post(
                f"/api/v1/projects/{self.project_id}/planning/blueprint/confirm",
                json={"revision": candidate["blueprint_revision"], "artifact_id": candidate["blueprint_artifact_id"]},
                headers=self.headers(),
            )
            self.assertEqual(200, confirmed.status_code, confirmed.text)

    def test_fact_publish_tolerates_outline_revision_drift(self) -> None:
        service = self.app.state.planning_service
        current = service.get(self.first, self.project_id)
        rows = current["literature_matrix"]["rows"]
        source_payload = {
            "source_matrix_artifact_id": current["matrix_artifact_id"],
            "expected_matrix_revision": current["matrix_revision"],
            "papers": [
                {
                    "paper_id": row["paper_id"],
                    "source_fingerprint": f"fact-{row['paper_id']}",
                    "index_summary": {},
                    "evidence_candidates": [],
                }
                for row in rows
            ],
        }
        built = {
            "papers": [
                {
                    "paper_id": row["paper_id"],
                    "status": "failed",
                    "facts": [],
                    "failed_fields": ["all"],
                    "error": "no source-addressable evidence",
                }
                for row in rows
            ]
        }

        with TestClient(self.app) as client:
            self.choose_outline(client, "reaction")
            selected = self.planning(client)

        self.assertEqual(current["matrix_artifact_id"], selected["matrix_artifact_id"])
        self.assertGreater(selected["matrix_revision"], current["matrix_revision"])
        published = service.publish_matrix_enrichment(
            self.first, self.project_id, source_payload, built
        )
        reloaded = service.get(self.first, self.project_id)

        self.assertGreater(published["matrix_revision"], selected["matrix_revision"])
        self.assertEqual(35, reloaded["matrix_enrichment"]["counts"]["failed"])
        self.assertEqual(0, reloaded["matrix_enrichment"]["counts"]["pending"])
        self.assertTrue(reloaded["outline_current"])

    def test_candidate_only_fact_enrichment_does_not_move_current_matrix(self) -> None:
        service = self.app.state.planning_service
        current = service.get(self.first, self.project_id)
        rows = current["literature_matrix"]["rows"]
        source_payload = {
            "source_matrix_artifact_id": current["matrix_artifact_id"],
            "expected_matrix_revision": current["matrix_revision"],
            "papers": [
                {
                    "paper_id": row["paper_id"],
                    "source_fingerprint": f"candidate-{row['paper_id']}",
                    "index_summary": {},
                    "evidence_candidates": [],
                }
                for row in rows
            ],
        }
        built = {
            "papers": [
                {
                    "paper_id": row["paper_id"],
                    "status": "failed",
                    "facts": [],
                    "failed_fields": ["all"],
                    "error": "no source-addressable evidence",
                }
                for row in rows
            ]
        }

        candidate = service.publish_matrix_enrichment(
            self.first,
            self.project_id,
            source_payload,
            built,
            candidate_only=True,
        )
        reloaded = service.get(self.first, self.project_id)

        self.assertEqual(current["matrix_artifact_id"], reloaded["matrix_artifact_id"])
        self.assertEqual(current["matrix_revision"], reloaded["matrix_revision"])
        self.assertEqual(len(rows), len(candidate["matrix_snapshot"]["rows"]))
        self.assertTrue(all(
            row["fact_enrichment"]["status"] == "failed"
            for row in candidate["matrix_snapshot"]["rows"]
        ))

    def test_fact_publish_uses_shared_excerpt_rules_and_required_roles(self) -> None:
        service = self.app.state.planning_service
        current = service.get(self.first, self.project_id)
        rows = current["literature_matrix"]["rows"]
        evidence_key = "sha256:shared-validator"
        source_payload = {
            "source_matrix_artifact_id": current["matrix_artifact_id"],
            "expected_matrix_revision": current["matrix_revision"],
            "classification_axes": [],
            "papers": [
                {
                    "paper_id": row["paper_id"],
                    "source_fingerprint": f"shared-{row['paper_id']}",
                    "index_summary": {},
                    "evidence_candidates": (
                        [
                            {
                                "evidence_key": evidence_key,
                                "content": (
                                    r"Terminal alkynes—including 1a–1c—and aldehydes "
                                    r"were combined at $25\,\mathrm { C }$."
                                ),
                                "content_type": "text",
                                "question_ids": ["method_conditions"],
                            }
                        ]
                        if row["paper_id"] == "P001"
                        else []
                    ),
                }
                for row in rows
            ],
        }
        built = {
            "papers": [
                (
                    {
                        "paper_id": row["paper_id"],
                        "status": "partial",
                        "facts": [
                            {
                                "fact_id": "MF-SHARED",
                                "fact_schema_version": "scientific-fact/2",
                                "field_id": "object_input",
                                "value": "The reported inputs were terminal alkynes and aldehydes.",
                                "support_excerpt": (
                                    r"Terminal alkynes-including 1a-1c-and aldehydes "
                                    r"were combined at $25\,\mathrm { C}$."
                                ),
                                "epistemic_status": "direct_source_report",
                                "confidence": 0.95,
                                "support_level": "direct",
                                "evidence_ceiling": "Only the stated inputs are supported.",
                                "assertion_ceiling": "direct_source_report",
                                "evidence_refs": [{"evidence_key": evidence_key}],
                            }
                        ],
                        "failed_fields": [],
                    }
                    if row["paper_id"] == "P001"
                    else {
                        "paper_id": row["paper_id"],
                        "status": "failed",
                        "facts": [],
                        "failed_fields": ["all"],
                    }
                )
                for row in rows
            ]
        }

        service.publish_matrix_enrichment(
            self.first,
            self.project_id,
            source_payload,
            built,
        )
        published = service.get(self.first, self.project_id)["literature_matrix"]
        paper = next(row for row in published["rows"] if row["paper_id"] == "P001")

        self.assertEqual("object_input", paper["scientific_facts"][0]["field_id"])
        self.assertEqual(
            ["object_input"],
            paper["fact_enrichment"]["supported_fact_roles"],
        )

    @staticmethod
    def isolated_reference_analysis(
        _principal,
        _project_id,
        *,
        candidate_id,
        safe_name,
        raw,
        matrix,
    ) -> dict:
        del safe_name, raw
        paper_ids = [row["paper_id"] for row in matrix["rows"]]
        representatives = paper_ids[:6]
        outline = (
            "# Selected Outline\n\n"
            "Scientific content source: current literature Matrix only.\n\n"
            "## Introduction and scope\n"
            f"Assigned papers: {', '.join(representatives)}.\n"
            "Purpose: define the current review scope.\n\n"
            "## 1. Copper allenation evidence\n"
            f"Assigned papers: {', '.join(paper_ids)}.\n"
            "Purpose: compare evidence from the current Matrix.\n\n"
            "## Conclusion and outlook\n"
            f"Assigned papers: {', '.join(representatives)}.\n"
            "Purpose: synthesize limitations and future directions.\n"
        )
        return {
            "candidate_id": candidate_id,
            "analysis_mode": "ai_style_only_transfer_v2",
            "content_source": "current_matrix_only",
            "reference_content_reused": False,
            "content_firewall": {
                "transfer_received_reference_text": False,
                "all_heading_levels_content_source": "current_matrix_only",
            },
            "reference_structure_metrics": {"heading_count": 3},
            "writing_style": {"organization_pattern": "progressive comparison"},
            "outline_md": outline,
        }

    def test_matrix_contains_entire_confirmed_selection(self) -> None:
        with TestClient(self.app) as client:
            payload = self.planning(client)
        self.assertEqual(35, len(payload["literature_matrix"]["rows"]))
        self.assertEqual(35, payload["matrix_sync"]["selected_paper_count"])
        self.assertNotIn("selection_fingerprint", payload["discovery_selection"])

    def test_reconfirmation_replaces_matrix_selection(self) -> None:
        seed_verified_matrix(self.app.state.planning_service, self.first, self.project_id)
        with TestClient(self.app) as client:
            self.choose_outline(client)
            blueprint_revision = self.app.state.workflow_repository.get_stage_state(
                self.first.user_id, self.project_id, "blueprint"
            )
            if blueprint_revision is None:
                response = self.generate_candidate(client,
                    json={"revision": 0},
                    headers=self.headers(),
                )
                self.assertEqual(200, response.status_code, response.text)
                confirmed_blueprint = client.post(f"/api/v1/projects/{self.project_id}/planning/blueprint/confirm",
                    json={"revision": 0, "artifact_id": response.json()["blueprint_artifact_id"]}, headers=self.headers())
                self.assertEqual(200, confirmed_blueprint.status_code, confirmed_blueprint.text)
            previous_blueprint = (
                self.app.state.workflow_repository.get_current_artifact(
                    self.first.user_id,
                    self.project_id,
                    "blueprint/section_blueprint.json",
                )
            )
            review = client.get(f"/api/v1/projects/{self.project_id}/discovery").json()
            for row in review["results"][0]["local_results"]:
                row["selected_for_matrix"] = row["paper_id"] in {"P001", "P035"}
            saved = client.put(
                f"/api/v1/projects/{self.project_id}/discovery",
                json={"revision": review["revision"], "results": review["results"]},
                headers=self.headers(),
            ).json()
            confirmed = client.post(
                f"/api/v1/projects/{self.project_id}/discovery/confirm",
                json={"revision": saved["revision"]},
                headers=self.headers(),
            )
            self.assertEqual(200, confirmed.status_code, confirmed.text)
            planning = self.planning(client)
        self.assertEqual(["P001", "P035"], [row["paper_id"] for row in planning["literature_matrix"]["rows"]])
        self.assertEqual(
            previous_blueprint.id,
            self.app.state.workflow_repository.get_current_artifact(
                self.first.user_id, self.project_id, "blueprint/section_blueprint.json"
            ).id,
        )
        self.assertFalse(planning["outline_current"])
        self.assertFalse(planning["blueprint_current"])
        self.assertEqual(
            "stale",
            self.app.state.workflow_repository.get_stage_state(
                self.first.user_id, self.project_id, "blueprint"
            ).status,
        )

    def test_matrix_row_edit_uses_revision(self) -> None:
        with TestClient(self.app) as client:
            current = self.planning(client)
            response = client.put(
                f"/api/v1/projects/{self.project_id}/planning/matrix/P001",
                json={
                    "revision": current["matrix_revision"],
                    "main_content": "A" * 320,
                    "mark_complete": True,
                },
                headers=self.headers(),
            )
            self.assertEqual(200, response.status_code, response.text)
            stale = client.put(
                f"/api/v1/projects/{self.project_id}/planning/matrix/P002",
                json={"revision": current["matrix_revision"], "main_content": "changed"},
                headers=self.headers(),
            )
            reloaded = self.planning(client)
        self.assertEqual(409, stale.status_code, stale.text)
        rows = {row["paper_id"]: row for row in reloaded["literature_matrix"]["rows"]}
        self.assertEqual("full_reading_complete", rows["P001"]["matrix_status"])
        self.assertEqual("", rows["P002"]["main_content"])

    def test_builtin_outline_loads_editable_content(self) -> None:
        with TestClient(self.app) as client:
            selected = self.choose_outline(client, "reaction")
        self.assertIn("##", selected["selected_outline_md"])
        self.assertIn("## Introduction\nSection role: introduction", selected["selected_outline_md"])
        self.assertNotIn(
            "Section role: conclusion",
            selected["selected_outline_md"],
        )
        self.assertIn("## 1. Cross-coupling", selected["selected_outline_md"])
        self.assertIn("## 2. Addition reactions", selected["selected_outline_md"])
        self.assertIn("## 3. Cyclization and annulation", selected["selected_outline_md"])
        self.assertTrue(selected["outline_complete"])
        self.assertEqual("reaction", selected["outline_style"])

    def test_topic_outline_intent_uses_query_plan_and_explicit_partitions(self) -> None:
        topic = (
            "Please write a review on allenation-of-terminal-alkynes (ATA). "
            "Focus on mono-, 1,3-di-, and trisubstituted allenes. "
            "Organize the review by reaction type and catalytic/promoting system "
            "(Cu, Zn, Cd, Ti, etc.), and separately discuss racemic ATA and "
            "enantioselective ATA (EATA)."
        )
        discovery = {
            "query_plan": {
                "group_by": ["reaction_type", "catalyst_or_method"],
            }
        }

        intent = _topic_outline_intent(topic, discovery)

        self.assertTrue(intent["available"])
        self.assertEqual("reaction_type", intent["primary_axis"])
        self.assertEqual(["catalyst_or_method"], intent["secondary_axes"])
        self.assertEqual(
            ["racemic ATA", "enantioselective ATA (EATA)"],
            intent["partitions"],
        )
        self.assertEqual(intent["partitions"], intent["required_partitions"])
        self.assertEqual(["Cu", "Zn", "Cd", "Ti"], intent["named_systems"])
        self.assertEqual(intent["named_systems"], intent["comparison_dimensions"])
        self.assertEqual(
            [
                "mono-, 1,3-di-, and trisubstituted allenes",
            ],
            intent["focus_dimensions"],
        )
        self.assertEqual(intent["focus_dimensions"], intent["outcome_dimensions"])
        self.assertEqual(
            "source_bounded_model_or_section_contract",
            intent["partition_trace_policy"],
        )
        self.assertEqual(
            "reaction_type",
            intent["classification_contract"]["primary_axis_id"],
        )
        self.assertEqual(
            ["reaction_type", "catalyst_or_method"],
            [
                axis["axis_id"]
                for axis in intent["classification_contract"]["axes"][:2]
            ],
        )
        self.assertEqual(64, len(intent["classification_contract"]["fingerprint"]))

    def test_topic_intent_keeps_comparison_examples_out_of_required_partitions(self) -> None:
        topic = (
            "Organize the review by intervention type and compare age groups "
            "(children, adults, older adults), then separately discuss randomized "
            "evidence and observational evidence."
        )

        intent = _topic_outline_intent(topic, None)

        self.assertEqual(
            ["randomized evidence", "observational evidence"],
            intent["required_partitions"],
        )
        self.assertNotIn("children", intent["required_partitions"])
        self.assertNotIn("adults", intent["required_partitions"])

    def test_topic_intent_splits_repaired_stereochemistry_from_reaction_hierarchy(self) -> None:
        intent = _topic_outline_intent(
            (
                "Organize the review by reaction type and separately discuss "
                "racemic ATA and enantioselective ATA."
            ),
            {"query_plan": {"group_by": ["reaction_type"]}},
            [
                {
                    "axis_id": "stereochemical_regime",
                    "label": "Stereochemical regime",
                    "source_surface": (
                        "Organize the review by reaction type and separately "
                        "discuss racemic ATA and enantioselective ATA"
                    ),
                    "source_type": "explicit_topic",
                    "axis_role": "primary_organization",
                    "heading_requirement": "primary_heading",
                    "semantic_repair": {"status": "auto_repaired"},
                    "partitions": [
                        {"partition_id": "racemic", "label": "racemic ATA"},
                        {
                            "partition_id": "enantioselective",
                            "label": "enantioselective ATA",
                        },
                    ],
                }
            ],
        )

        self.assertEqual("reaction_type", intent["primary_axis"])
        self.assertEqual(
            ["stereochemical_regime"], intent["secondary_axes"]
        )
        self.assertEqual(
            ["racemic ATA", "enantioselective ATA"],
            intent["required_partitions"],
        )
        self.assertEqual(
            "required_independent_discussion",
            intent["classification_axes"][1]["axis_role"],
        )
        self.assertEqual(
            intent["classification_contract"]["axes"],
            intent["classification_axes"],
        )

    def test_topic_intent_reads_examples_for_other_chemistry_axes(self) -> None:
        topic = (
            "Review the transformation categorized by substrates "
            "(aryl halides, alkenes, organoboron reagents, etc.)."
        )
        discovery = {"query_plan": {"group_by": ["substrate"]}}

        intent = _topic_outline_intent(topic, discovery)

        self.assertEqual("substrate", intent["primary_axis"])
        self.assertEqual(
            ["aryl halides", "alkenes", "organoboron reagents"],
            intent["axis_examples"]["substrate"],
        )
        self.assertEqual([], intent["required_partitions"])

    def test_topic_guided_outline_supports_product_as_primary_axis(self) -> None:
        service = self.app.state.planning_service
        rows = [{"paper_id": "P001"}, {"paper_id": "P002"}]
        intent = {
            "available": True,
            "primary_axis": "product",
            "secondary_axes": ["catalyst_or_method"],
            "required_partitions": [],
        }

        outline = service._topic_outline_document(
            rows,
            tags_by_paper={
                "P001": {"product": "heterocycles"},
                "P002": {"product": "pharmaceutical compounds"},
            },
            text_by_paper={
                "P001": "Organocatalysis furnished a heterocyclic compound.",
                "P002": "Enzymatic methods furnished a pharmaceutical product.",
            },
            taxonomy_profile="chemistry_general",
            intent=intent,
        )

        self.assertIn("Primary structure: Topic-guided (product class", outline)
        self.assertIn("Heterocycles", outline)
        self.assertIn("Pharmaceutical compounds", outline)

    def test_topic_partition_matching_uses_declared_terms_without_domain_defaults(self) -> None:
        partitions = ["photochemical conditions", "electrochemical conditions"]

        self.assertEqual(
            "photochemical conditions",
            _topic_partition_for_text(
                "The reaction was performed under photochemical conditions.",
                partitions,
            ),
        )
        self.assertEqual(
            "",
            _topic_partition_for_text(
                "The source reports thermal activation but no requested partition.",
                partitions,
            ),
        )

    def test_topic_guided_outline_keeps_partitions_inside_primary_axis(self) -> None:
        service = self.app.state.planning_service
        rows = [{"paper_id": f"P{index:03d}"} for index in range(1, 5)]
        tags_by_paper = {
            "P001": {"reaction_type": "three-component coupling"},
            "P002": {"reaction_type": "three-component coupling"},
            "P003": {"reaction_type": "homologation"},
            "P004": {"reaction_type": "homologation"},
        }
        text_by_paper = {
            "P001": "Racemic Cu-promoted terminal alkyne allenation afforded an allene.",
            "P002": "Enantioselective Cu-catalyzed ATA afforded 95% ee.",
            "P003": "Racemic zinc-promoted homologation of a terminal alkyne.",
            "P004": "Asymmetric homologation gave an enantioenriched allene.",
        }
        intent = {
            "available": True,
            "primary_axis": "reaction_type",
            "secondary_axes": ["catalyst_or_method"],
            "required_partitions": ["racemic ATA", "enantioselective ATA (EATA)"],
            "comparison_dimensions": ["Cu", "Zn", "Cd", "Ti"],
            "focus_dimensions": [
                "monosubstituted allenes",
                "1,3-disubstituted allenes",
                "trisubstituted allenes",
            ],
        }

        outline = service._topic_outline_document(
            rows,
            tags_by_paper=tags_by_paper,
            text_by_paper=text_by_paper,
            taxonomy_profile="allene",
            intent=intent,
        )

        self.assertIn("## 1. Three-component coupling", outline)
        self.assertIn("## 2. Homologation", outline)
        self.assertNotIn("Racemic ATA — Three-component coupling", outline)
        self.assertIn("Assigned papers: P001, P002.", outline)
        self.assertIn(
            "Topic-requested independent discussion represented by assigned evidence: racemic ATA, enantioselective ATA (EATA).",
            outline,
        )
        self.assertIn("catalytic or promoting system", outline)
        self.assertIn(
            "Focus dimensions: monosubstituted allenes, "
            "1,3-disubstituted allenes, trisubstituted allenes.",
            outline,
        )

    def test_topic_guided_outline_prefers_evidence_bound_model_partition(self) -> None:
        service = self.app.state.planning_service
        rows = [
            {
                "paper_id": "P001",
                "topic_partition_classification": {
                    "status": "classified",
                    "partition": "randomized evidence",
                    "confidence": 0.92,
                    "evidence_refs": [{"evidence_key": "sha256:source"}],
                },
            }
        ]
        outline = service._topic_outline_document(
            rows,
            tags_by_paper={"P001": {"reaction_type": "controlled comparison"}},
            text_by_paper={
                "P001": "The title and abstract use no literal partition label."
            },
            taxonomy_profile="general",
            intent={
                "available": True,
                "primary_axis": "reaction_type",
                "secondary_axes": [],
                "required_partitions": [
                    "randomized evidence",
                    "observational evidence",
                ],
            },
        )

        self.assertIn("## 1. Controlled comparison", outline)
        self.assertNotIn("Randomized evidence — Controlled comparison", outline)
        self.assertIn(
            "Topic-requested independent discussion represented by assigned evidence: randomized evidence.",
            outline,
        )

    def test_completed_model_boundary_is_not_overruled_by_keyword_match(self) -> None:
        routed = _topic_partition_for_row(
            {
                "topic_partition_classification": {
                    "status": "boundary",
                    "partition": "",
                    "confidence": 0.42,
                    "evidence_refs": [],
                }
            },
            ["randomized evidence", "observational evidence"],
            "The related-work paragraph mentions randomized evidence.",
        )

        self.assertEqual("", routed)

    def test_blueprint_partition_routes_follow_source_bound_matrix_classification(self) -> None:
        sections = [
            {
                "section_id": "S02",
                "title": "Controlled comparisons",
                "section_role": "body",
                "primary_papers": ["P001", "P002"],
            }
        ]
        rows = {
            "P001": {
                "paper_id": "P001",
                "topic_partition_classification": {
                    "status": "classified",
                    "partition": "randomized evidence",
                    "confidence": 0.91,
                    "evidence_refs": [{"evidence_key": "sha256:randomized"}],
                },
            },
            "P002": {
                "paper_id": "P002",
                "topic_partition_classification": {
                    "status": "boundary",
                    "partition": "",
                    "confidence": 0.45,
                    "evidence_refs": [],
                },
            },
        }

        routes, support = _topic_partition_routes(
            sections,
            rows,
            ["randomized evidence", "observational evidence"],
            {
                "P001": "The literal label is intentionally absent.",
                "P002": "Observational evidence appears only in related work.",
            },
        )

        self.assertEqual(
            {"randomized evidence": ["P001"]}, routes["S02"]
        )
        self.assertEqual(["P001"], support["randomized evidence"])
        self.assertEqual([], support["observational evidence"])

    def test_required_partitions_upgrade_legacy_outline_fields(self) -> None:
        partitions = _required_topic_partitions_from_outline(
            {
                "classification_basis": {
                    "required_outline_partitions": [
                        "randomized evidence",
                        "observational evidence",
                    ]
                },
                "classification_contract": {
                    "axes": [
                        {
                            "axis_role": "required_independent_discussion",
                            "partitions": [
                                {"label": "randomized evidence"},
                                {"label": "qualitative evidence"},
                            ],
                        }
                    ]
                },
            }
        )

        self.assertEqual(
            [
                "randomized evidence",
                "observational evidence",
                "qualitative evidence",
            ],
            partitions,
        )

    def test_formal_matrix_tag_routes_before_legacy_topic_text_match(self) -> None:
        routed = _topic_partition_for_row(
            {
                "evidence_backed_tags": {
                    "study_design": [
                        {
                            "partition_label": "randomized evidence",
                            "confidence": 0.93,
                            "fact_ids": ["MF-1"],
                            "evidence_refs": [{"evidence_key": "sha256:formal"}],
                        }
                    ]
                }
            },
            ["randomized evidence", "observational evidence"],
            "The background mentions observational evidence.",
        )
        self.assertEqual("randomized evidence", routed)

    def test_topic_guided_outline_uses_configured_taxonomy_without_topic_branch(self) -> None:
        service = self.app.state.planning_service
        rows = [{"paper_id": f"P{index:03d}"} for index in range(1, 4)]
        text_by_paper = {
            "P001": "Transition metal catalysis enabled a cross-coupling reaction.",
            "P002": "Visible light photochemical conditions enabled a cycloaddition.",
            "P003": "An electrochemical oxidation furnished the target product.",
        }
        intent = {
            "available": True,
            "primary_axis": "reaction_type",
            "secondary_axes": ["catalyst_or_method"],
            "partitions": [],
            "comparison_dimensions": ["operating conditions"],
        }

        outline = service._topic_outline_document(
            rows,
            tags_by_paper={},
            text_by_paper=text_by_paper,
            taxonomy_profile="chemistry_general",
            intent=intent,
        )

        self.assertIn("Cross-coupling", outline)
        self.assertIn("Cyclization and annulation", outline)
        self.assertIn("Oxidation and reduction", outline)
        self.assertIn("transition-metal catalysis", outline)
        self.assertIn("photochemical methods", outline)
        self.assertIn("electrochemical methods", outline)

    def test_topic_guided_outline_retains_unresolved_primary_study_as_boundary(self) -> None:
        service = self.app.state.planning_service
        outline = service._topic_outline_document(
            [{"paper_id": "P001"}],
            tags_by_paper={"P001": {}},
            text_by_paper={
                "P001": "A selected primary experiment whose current evidence has no declared route."
            },
            taxonomy_profile="general_academic",
            intent={
                "available": True,
                "primary_axis": "reaction_type",
                "secondary_axes": [],
                "required_partitions": [],
            },
        )

        self.assertIn("## 1. Cross-category comparison", outline)
        self.assertIn("Assigned papers: P001.", outline)
        self.assertIn("Boundary rationale:", outline)

    def test_reselecting_current_outline_is_idempotent(self) -> None:
        with TestClient(self.app) as client:
            selected = self.choose_outline(client, "reaction")
            repeated = client.put(
                f"/api/v1/projects/{self.project_id}/planning/outline",
                json={
                    "revision": selected["matrix_revision"],
                    "outline_style": "reaction",
                },
                headers=self.headers(),
            )
        self.assertEqual(200, repeated.status_code, repeated.text)
        payload = repeated.json()
        self.assertTrue(payload["unchanged"])
        self.assertEqual(selected["matrix_revision"], payload["matrix_revision"])
        self.assertEqual(selected["outline_artifact_id"], payload["outline_artifact_id"])

    def test_builtin_outline_styles_use_distinct_metadata_axes(self) -> None:
        with TestClient(self.app) as client:
            substrate = self.choose_outline(client, "substrate")["selected_outline_md"]
            catalyst = self.choose_outline(client, "catalyst")["selected_outline_md"]
            reaction = self.choose_outline(client, "reaction")["selected_outline_md"]
        self.assertIn("## 1. Aromatic substrates", substrate)
        self.assertIn("## 1. Transition-metal catalysis", catalyst)
        self.assertIn("## 1. Cross-coupling", reaction)
        self.assertNotEqual(substrate, catalyst)
        self.assertNotEqual(catalyst, reaction)

    def test_chemistry_substrate_outline_can_group_common_allene_precursors(self) -> None:
        service = self.app.state.planning_service
        rows = [{"paper_id": f"P{index:03d}"} for index in range(1, 8)]
        text_by_paper = {
            "P001": "enantioselective conversion of a propargylic alcohol substrate",
            "P002": "scope of substituted propargyl alcohol starting materials",
            "P003": "catalytic synthesis from a terminal alkyne",
            "P004": "three-component reaction using 1-alkynes",
            "P005": "cyclization of a conjugated enyne",
            "P006": "asymmetric transformation of substituted 1,3-enynes",
            "P007": "terminal alkynes as allene precursors; the abstract later compares propargylic alcohol chemistry",
        }

        groups = service._lexical_outline_candidates(
            rows,
            text_by_paper,
            tag_key="substrate",
            taxonomy_profile="allene",
        )

        self.assertEqual(["P001", "P002"], groups["propargylic alcohols"])
        self.assertEqual(["P003", "P004", "P007"], groups["terminal alkynes"])
        self.assertEqual(["P005", "P006"], groups["enynes"])

    def test_paper_text_cannot_activate_allene_rules_without_topic_routing(self) -> None:
        groups = self.app.state.planning_service._lexical_outline_candidates(
            [{"paper_id": "P001"}],
            {"P001": "allene synthesis from a propargylic alcohol"},
            tag_key="substrate",
            taxonomy_profile="chemistry_general",
        )

        self.assertEqual({ROUTING_REQUIRED_LABEL: ["P001"]}, groups)

    def test_chemistry_outline_reroutes_allene_evidence_without_catch_all(self) -> None:
        service = self.app.state.planning_service
        rows = [{"paper_id": f"P{index:03d}"} for index in range(1, 7)]
        text_by_paper = {
            "P001": "enantioselective isomerization of 3-alkynoates to chiral allenoates",
            "P002": "this review will highlight allenes in catalytic asymmetric synthesis",
            "P003": "phase-transfer functionalization of 1-alkylallene-1,3-dicarboxylates",
            "P004": "palladium synthesis of axially chiral (allenylmethyl)silanes",
            "P005": "hydroboration of but-1-en-3-ynes to axially chiral allenylboranes",
            "P006": "resolution of an allene hydrocarbon into optical antipodes",
        }

        outline = service._outline_document(
            "substrate",
            rows,
            tags_by_paper={paper_id: {} for paper_id in text_by_paper},
            text_by_paper=text_by_paper,
            taxonomy_profile="allene",
        )

        self.assertNotIn("Routing required", outline)
        self.assertIn("Context papers: P002.", outline)
        self.assertIn("## 1. Alkynoates", outline)
        self.assertIn("Preformed substituted allenes", outline)
        self.assertIn("Silyl-substituted diene precursors", outline)
        self.assertIn("Enynes", outline)

    def test_chemistry_outline_routes_resolution_and_enynamide_titles(self) -> None:
        service = self.app.state.planning_service
        rows = [{"paper_id": "P001"}, {"paper_id": "P002"}]
        groups = service._lexical_outline_candidates(
            rows,
            {
                "P001": "chemoenzymatic dynamic kinetic resolution of axially chiral allenes",
                "P002": "rhodium-catalyzed 1,6-addition of arylboronic acids to enynamides",
            },
            tag_key="substrate",
            taxonomy_profile="allene",
        )

        self.assertEqual(["P001"], groups["preformed substituted allenes"])
        self.assertEqual(["P002"], groups["enynes"])
        self.assertNotIn("Routing required — reassign these papers", groups)

    def test_generated_routing_placeholder_is_repaired_before_blueprint(self) -> None:
        service = self.app.state.planning_service
        sections = [
            {
                "title": "Introduction",
                "section_role": "introduction",
                "paper_ids": [],
            },
            {
                "title": "preformed substituted allenes",
                "section_role": "body",
                "paper_ids": ["P003"],
            },
            {
                "title": "enynes",
                "section_role": "body",
                "paper_ids": ["P004"],
            },
            {
                "title": "Routing required — reassign these papers",
                "section_role": "body",
                "paper_ids": ["P001", "P002"],
            },
            {
                "title": "Conclusion",
                "section_role": "conclusion",
                "paper_ids": [],
            },
        ]
        repaired, adjustments = service._auto_repair_generated_routing_sections(
            sections,
            [{"paper_id": f"P00{index}"} for index in range(1, 5)],
            {
                "P001": "chemoenzymatic dynamic kinetic resolution of axially chiral allenes",
                "P002": "enantioselective 1,6-addition to enynamides",
            },
            outline_style="substrate",
            taxonomy_profile="allene",
        )

        by_title = {section["title"]: section for section in repaired}
        self.assertNotIn("Routing required — reassign these papers", by_title)
        self.assertEqual(
            ["P003", "P001"],
            by_title["preformed substituted allenes"]["paper_ids"],
        )
        self.assertEqual(["P004", "P002"], by_title["enynes"]["paper_ids"])
        self.assertEqual(2, len(adjustments))

    def test_generated_boundary_section_is_repaired_again_from_scientific_objects(self) -> None:
        service = self.app.state.planning_service
        repaired, adjustments = service._auto_repair_generated_routing_sections(
            [
                {"title": "Introduction", "section_role": "introduction", "paper_ids": []},
                {
                    "title": "Cross-category evidence and boundary cases",
                    "section_role": "body",
                    "paper_ids": ["P001", "P002"],
                },
                {"title": "Conclusion", "section_role": "conclusion", "paper_ids": []},
            ],
            [{"paper_id": "P001"}, {"paper_id": "P002"}],
            {
                "P001": "propargylic benzoates are converted to axially chiral allenes",
                "P002": "activated enynes are converted to axially chiral allenes",
            },
            outline_style="substrate",
            taxonomy_profile="allene",
        )

        by_title = {section["title"]: section for section in repaired}
        self.assertNotIn("Cross-category evidence and boundary cases", by_title)
        self.assertEqual(["P001"], by_title["activated propargylic derivatives"]["paper_ids"])
        self.assertEqual(["P002"], by_title["enynes"]["paper_ids"])
        self.assertEqual(2, len(adjustments))

    def test_unresolved_primary_study_is_not_relabelled_as_introduction_context(self) -> None:
        service = self.app.state.planning_service
        repaired, adjustments = service._auto_repair_generated_routing_sections(
            [
                {
                    "title": "Introduction",
                    "section_role": "introduction",
                    "paper_ids": [],
                    "context_paper_ids": [],
                },
                {
                    "title": "Routing required — reassign these papers",
                    "section_role": "body",
                    "paper_ids": ["P001"],
                },
                {
                    "title": "Conclusion",
                    "section_role": "conclusion",
                    "paper_ids": [],
                },
            ],
            [{"paper_id": "P001"}],
            {"P001": "A selected primary experiment with no supported route."},
            outline_style="reaction",
            taxonomy_profile="general_academic",
        )

        introduction = next(
            section
            for section in repaired
            if section["section_role"] == "introduction"
        )
        boundary = next(
            section
            for section in repaired
            if section.get("title") == "Cross-category comparison"
        )
        self.assertEqual([], introduction["context_paper_ids"])
        self.assertEqual(["P001"], boundary["paper_ids"])
        self.assertTrue(boundary["boundary_rationale"])
        self.assertEqual("unresolved_classification_retained", adjustments[0]["method"])
        self.assertEqual(["P001"], adjustments[0]["paper_ids"])

    def test_generated_coverage_reconciliation_uses_bounded_fact_excerpt(self) -> None:
        service = self.app.state.planning_service
        sections = [
            {
                "title": "Introduction",
                "section_role": "introduction",
                "paper_ids": [],
                "context_paper_ids": [],
            },
            {
                "title": "Formaldehyde-based terminal-alkyne homologation",
                "section_role": "body",
                "paper_ids": ["P001"],
            },
            {
                "title": "Conclusion",
                "section_role": "conclusion",
                "paper_ids": [],
            },
        ]
        repaired, adjustments = service._reconcile_generated_outline_coverage(
            sections,
            [{"paper_id": "P001"}, {"paper_id": "P002"}],
            {"P001": {}, "P002": {}},
            {
                "P001": "terminal alkyne homologation",
                "P002": (
                    "CuI and paraformaldehyde were combined with a terminal alkyne "
                    "to synthesize a terminal allene."
                ),
            },
            outline_style="topic-guided",
            taxonomy_profile="allene",
            tag_key_override="reaction_type",
            axis_label_override="reaction type",
        )

        target = next(
            section
            for section in repaired
            if section["title"] == "Formaldehyde-based terminal-alkyne homologation"
        )
        self.assertEqual(["P001", "P002"], target["paper_ids"])
        self.assertEqual(
            "coverage_reconciliation_evidence_route",
            adjustments[0]["method"],
        )

    def test_topic_guided_repair_preserves_non_default_primary_axis(self) -> None:
        service = self.app.state.planning_service
        repaired, _adjustments = service._auto_repair_generated_routing_sections(
            [
                {
                    "title": "Routing required — reassign these papers",
                    "section_role": "body",
                    "paper_ids": ["P001"],
                }
            ],
            [{"paper_id": "P001"}],
            {"P001": "The reaction furnished a heterocyclic compound."},
            outline_style="topic-guided",
            taxonomy_profile="chemistry_general",
            tag_key_override="product",
            axis_label_override="product class",
        )

        self.assertEqual("heterocycles", repaired[0]["title"])
        self.assertIn("product class", repaired[0]["purpose"])

    def test_account_language_is_contextual_evidence(self) -> None:
        contextual = self.app.state.planning_service._contextual_outline_paper_ids(
            [{"paper_id": "P001"}, {"paper_id": "P002"}],
            {"P001": {}, "P002": {}},
            {
                "P001": "The account concerns palladium-catalyzed cyclization reactions.",
                "P002": "We report a controlled primary study.",
            },
        )

        self.assertEqual(["P001"], contextual)

    def test_generated_body_is_not_realigned_from_unverified_object_words(self) -> None:
        repaired, adjustments = (
            self.app.state.planning_service._realign_generated_body_sections(
                [
                    {"title": "Introduction", "section_role": "introduction", "paper_ids": []},
                    {
                        "title": "allenoates",
                        "section_role": "body",
                        "paper_ids": ["P001"],
                    },
                    {"title": "Conclusion", "section_role": "conclusion", "paper_ids": []},
                ],
                [
                    {
                        "paper_id": "P001",
                        "scientific_facts": [
                            {
                                "field_id": "object_input",
                                "value": "nitroalkanes and activated enynes",
                            }
                        ],
                    }
                ],
                {
                    "P001": (
                        "nitroalkanes and activated enynes are the input objects. "
                        "enantioselective synthesis of axially chiral allenes and allenoates"
                    )
                },
                outline_style="substrate",
                taxonomy_profile="allene",
            )
        )

        by_title = {section["title"]: section for section in repaired}
        self.assertEqual(["P001"], by_title["allenoates"]["paper_ids"])
        self.assertEqual([], adjustments)

    def test_equivalent_generated_body_sections_merge_without_cross_category_guessing(self) -> None:
        repaired, adjustments = (
            self.app.state.planning_service._merge_equivalent_generated_body_sections(
                [
                    {
                        "title": "Method Family A",
                        "section_role": "body",
                        "paper_ids": ["P001"],
                    },
                    {
                        "title": "method-family A",
                        "section_role": "body",
                        "paper_ids": ["P002"],
                    },
                    {
                        "title": "Method Family B",
                        "section_role": "body",
                        "paper_ids": ["P003"],
                    },
                ]
            )
        )

        self.assertEqual(2, len(repaired))
        self.assertEqual(["P001", "P002"], repaired[0]["paper_ids"])
        self.assertEqual(["P003"], repaired[1]["paper_ids"])
        self.assertEqual(
            "equivalent_primary_axis_category_merge",
            adjustments[0]["method"],
        )

    def test_topic_partition_survives_scientific_object_realignment(self) -> None:
        repaired, _adjustments = (
            self.app.state.planning_service._realign_generated_body_sections(
                [
                    {
                        "title": "Randomized evidence — Allenoates",
                        "section_role": "body",
                        "topic_partition": "randomized evidence",
                        "paper_ids": ["P001"],
                    }
                ],
                [
                    {
                        "paper_id": "P001",
                        "scientific_facts": [
                            {
                                "field_id": "object_input",
                                "value": "activated enynes",
                            }
                        ],
                    }
                ],
                {
                    "P001": (
                        "activated enynes furnished an axially chiral allene product"
                    )
                },
                outline_style="topic-guided",
                taxonomy_profile="allene",
                tag_key_override="substrate",
                axis_label_override="substrate class",
            )
        )

        self.assertEqual(1, len(repaired))
        self.assertEqual("randomized evidence", repaired[0]["topic_partition"])
        self.assertEqual("Randomized evidence — Allenoates", repaired[0]["title"])

    def test_topic_boundary_rationale_survives_primary_axis_realignment(self) -> None:
        rationale = (
            "The source does not positively establish one requested Topic partition."
        )
        repaired, _adjustments = (
            self.app.state.planning_service._realign_generated_body_sections(
                [
                    {
                        "title": "Topic-partition boundary cases — ATA",
                        "section_role": "body",
                        "topic_partition": TOPIC_PARTITION_BOUNDARY_LABEL,
                        "boundary_rationale": rationale,
                        "paper_ids": ["P001"],
                    }
                ],
                [
                    {
                        "paper_id": "P001",
                        "scientific_facts": [
                            {
                                "field_id": "transformation",
                                "value": "terminal alkyne allenation",
                            }
                        ],
                    }
                ],
                {"P001": "terminal alkyne allenation furnished an allene"},
                outline_style="topic-guided",
                taxonomy_profile="allene",
                tag_key_override="reaction_type",
                axis_label_override="reaction type",
            )
        )

        self.assertEqual(1, len(repaired))
        self.assertEqual(
            TOPIC_PARTITION_BOUNDARY_LABEL,
            repaired[0]["topic_partition"],
        )
        self.assertEqual("Topic-partition boundary cases — ATA", repaired[0]["title"])
        self.assertEqual(rationale, repaired[0]["boundary_rationale"])

    def test_outline_sources_prefer_confirmed_and_ignore_legacy_automatic_tags(self) -> None:
        service = self.app.state.planning_service
        confirmed_tags, _ = service._outline_sources(
            self.first,
            [
                {
                    "paper_id": "P001",
                    "project_tag_review_status": "confirmed",
                    "project_tags": {
                        "reaction_type": ["project-specific transformation"]
                    },
                }
            ],
        )
        pending_tags, _ = service._outline_sources(
            self.first,
            [
                {
                    "paper_id": "P001",
                    "project_tag_review_status": "pending",
                    "project_tags": {"reaction_type": ["unreviewed suggestion"]},
                }
            ],
        )
        automatic_tags, _ = service._outline_sources(
            self.first,
            [
                {
                    "paper_id": "P001",
                    "project_tag_review_status": "automatic",
                    "project_tags": {
                        "reaction_type": ["automatically assessed transformation"]
                    },
                }
            ],
        )
        self.assertEqual(
            ["project-specific transformation"],
            confirmed_tags["P001"]["reaction_type"],
        )
        self.assertEqual("cross-coupling", pending_tags["P001"]["reaction_type"])
        self.assertEqual("cross-coupling", automatic_tags["P001"]["reaction_type"])

    def test_outline_sources_do_not_promote_unreviewed_agent_routing(self) -> None:
        with patch(
            "review_writer_api.domain_services.planning.verified_structured_tags",
            return_value={},
        ):
            tags, _ = self.app.state.planning_service._outline_sources(
                self.first,
                [
                    {
                        "paper_id": "P001",
                        "routing_recommendation": {
                            "axis_id": "reaction_type",
                            "status": "classified",
                            "label": "aldehyde-based three-component ATA",
                            "confidence": 0.96,
                            "evidence_refs": [{"evidence_key": "sha256:route"}],
                        },
                    }
                ],
            )

        self.assertNotIn("reaction_type", tags["P001"])

    def test_outline_sources_include_bounded_fact_excerpt_for_routing(self) -> None:
        with patch(
            "review_writer_api.domain_services.planning.verified_structured_tags",
            return_value={},
        ):
            _tags, text_by_paper = self.app.state.planning_service._outline_sources(
                self.first,
                [
                    {
                        "paper_id": "P001",
                        "scientific_facts": [
                            {
                                "field_id": "catalyst_or_method",
                                "value": "Copper-based system",
                                "support_level": "direct",
                                "evidence_refs": [{"evidence_key": "registered"}],
                                "support_excerpt": (
                                    "CuI and paraformaldehyde were added to the terminal "
                                    "alkyne substrate."
                                ),
                            }
                        ],
                    }
                ],
            )

        self.assertIn("paraformaldehyde", text_by_paper["P001"])
        self.assertIn("terminal alkyne substrate", text_by_paper["P001"])

    def test_outline_sources_ignore_unverified_library_tags(self) -> None:
        with self.sessions.begin() as session:
            paper = session.scalar(
                select(LibraryPaper).where(LibraryPaper.paper_id == "P001")
            )
            metadata = dict(paper.metadata_json)
            structured = dict(metadata["structured_tags"])
            structured["human_checked"] = False
            metadata["structured_tags"] = structured
            paper.metadata_json = metadata

        tags, _ = self.app.state.planning_service._outline_sources(
            self.first,
            [{"paper_id": "P001", "project_tag_review_status": "pending"}],
        )

        self.assertEqual({}, tags["P001"])

    def test_custom_outline_starts_blank(self) -> None:
        with TestClient(self.app) as client:
            selected = self.choose_outline(client, "custom")
        self.assertEqual("", selected["selected_outline_md"])
        self.assertFalse(selected["outline_complete"])

    def test_whole_outline_recommendation_uses_evidence_and_never_falls_back_to_first_papers(self) -> None:
        service = self.app.state.planning_service
        original_gateway = service.model_gateway
        service.model_gateway = None
        try:
            with patch.object(
                service,
                "_outline_sources",
                return_value=(
                    {},
                    {
                        "P001": "ketone ata allenylation evidence",
                        "P002": "aldehyde ata three component evidence",
                    },
                ),
            ), TestClient(self.app) as client:
                current = self.planning(client)
                response = client.post(
                    f"/api/v1/projects/{self.project_id}/planning/outline/recommendations",
                    json={
                        "revision": current["matrix_revision"],
                        "outline_md": (
                            "## 1. Ketone ATA\n"
                            "Purpose: Compare ketone reactions.\n\n"
                            "## 2. Aldehyde ATA\n"
                            "Purpose: Compare aldehyde three-component reactions.\n"
                        ),
                    },
                    headers=self.headers(),
                )
            self.assertEqual(200, response.status_code, response.text)
            payload = response.json()
            by_section = {
                item["section_index"]: item["paper_ids"] for item in payload["sections"]
            }
            self.assertEqual(["P001"], by_section[0])
            self.assertEqual(["P002"], by_section[1])
            self.assertEqual(2, payload["summary"]["recommended_paper_count"])
            self.assertEqual(33, payload["summary"]["unassigned_paper_count"])
        finally:
            service.model_gateway = original_gateway

    def test_whole_outline_recommendation_uses_one_model_pass_for_ambiguous_papers(self) -> None:
        service = self.app.state.planning_service
        original_gateway = service.model_gateway
        gateway = SimpleNamespace(
            environment_for_job=lambda _job: ({}, {"REVIEW_WRITER_TASK_TOKEN": "token"}),
            complete=AsyncMock(
                return_value={
                    "output_text": json.dumps(
                        {
                            "assignments": [
                                {
                                    "paper_id": "P001",
                                    "section_index": 1,
                                    "confidence": 0.9,
                                    "reason": "The extracted object matches aldehyde ATA.",
                                }
                            ]
                        }
                    )
                }
            ),
        )
        service.model_gateway = gateway
        try:
            with (
                patch.object(service, "_outline_sources", return_value=({}, {})),
                patch.object(
                    service,
                    "_begin_gateway_job",
                    return_value=SimpleNamespace(job_id=str(uuid.uuid4())),
                ),
                patch.object(service, "_finish_gateway_job"),
                TestClient(self.app) as client,
            ):
                current = self.planning(client)
                response = client.post(
                    f"/api/v1/projects/{self.project_id}/planning/outline/recommendations",
                    json={
                        "revision": current["matrix_revision"],
                        "outline_md": (
                            "## 1. Ketone ATA\nPurpose: Compare ketone reactions.\n\n"
                            "## 2. Aldehyde ATA\nPurpose: Compare aldehyde reactions.\n"
                        ),
                    },
                    headers=self.headers(),
                )
            self.assertEqual(200, response.status_code, response.text)
            self.assertEqual(["P001"], response.json()["sections"][1]["paper_ids"])
            self.assertEqual(1, response.json()["summary"]["model_resolved_paper_count"])
            gateway.complete.assert_awaited_once()
        finally:
            service.model_gateway = original_gateway

    def test_manual_outline_save_versions_content(self) -> None:
        with TestClient(self.app) as client:
            selected = self.choose_outline(client, "custom")
            outline = "# Review\n\n## 1. Introduction\nAssigned papers: P001, P002.\nPurpose: scope.\n"
            response = client.put(
                f"/api/v1/projects/{self.project_id}/planning/outline",
                json={
                    "revision": selected["matrix_revision"],
                    "outline_style": "custom",
                    "outline_md": outline,
                },
                headers=self.headers(),
            )
            self.assertEqual(200, response.status_code, response.text)
            saved = response.json()
            with patch.object(self.app.state.planning_service, "_publish_files") as publish:
                stale = client.put(
                    f"/api/v1/projects/{self.project_id}/planning/outline",
                    json={
                        "revision": selected["matrix_revision"],
                        "outline_style": "custom",
                        "outline_md": outline.replace("scope", "stale"),
                    },
                    headers=self.headers(),
                )
                publish.assert_not_called()
            reloaded = self.planning(client)
        self.assertEqual(409, stale.status_code)
        self.assertEqual(saved["outline_artifact_id"], reloaded["outline_selection"]["artifact_id"])
        self.assertIn("Purpose: scope.", reloaded["selected_outline_md"])

    def test_saved_outline_appears_in_comparison(self) -> None:
        with TestClient(self.app) as client:
            self.choose_outline(client, "catalyst")
            payload = self.planning(client)
        candidates = {item["candidate_id"]: item for item in payload["outline_candidates"]}
        self.assertIn("saved-current", candidates)
        self.assertEqual(payload["selected_outline_md"], candidates["saved-current"]["outline_md"])

    def test_empty_sections_survive_outline_blueprint_and_writing_handoff(self) -> None:
        all_papers = ", ".join(f"P{index:03d}" for index in range(1, 36))
        outline = (
            "## Introduction\nPurpose: Frame the review.\n\n"
            "## Historical development\nPurpose: Explain the field's development.\n\n"
            f"## Catalyst comparison\nAssigned papers: {all_papers}.\n\n"
            "## Conclusion\nPurpose: Synthesize body evidence.\n"
        )
        with TestClient(self.app) as client:
            selected = self.choose_outline(client, "custom")
            saved = client.put(
                f"/api/v1/projects/{self.project_id}/planning/outline",
                json={"revision": selected["matrix_revision"], "outline_style": "custom", "outline_md": outline},
                headers=self.headers(),
            )
            self.assertEqual(200, saved.status_code, saved.text)
            self.assertIn("## Historical development", self.planning(client)["selected_outline_md"])
            generated = self.generate_candidate(client,
                json={"revision": 0},
                headers=self.headers(),
            )
            self.assertEqual(200, generated.status_code, generated.text)
            blueprint = generated.json()["section_blueprint"]
            self.assertEqual(saved.json()["outline_artifact_id"], blueprint["candidate_base_outline_artifact_id"])
            history = next(section for section in blueprint["sections"] if section["title"] == "Historical development")
            self.assertEqual([], history["primary_papers"])
            self.assertEqual([], history["scientific_claims"])
            self.assertEqual("not_reviewed", history["evidence_readiness"]["status"])
            for section in blueprint["sections"]:
                if section["section_role"] in {"introduction", "conclusion"}:
                    self.assertEqual("synthesis", section["evidence_readiness"]["status"])
                    self.assertTrue(section["supporting_papers"])
            confirmed = client.post(
                f"/api/v1/projects/{self.project_id}/planning/blueprint/confirm",
                json={"revision": generated.json()["blueprint_revision"]},
                headers=self.headers(),
            )
            self.assertEqual(200, confirmed.status_code, confirmed.text)
        self.assertTrue(history["generation_eligible"])
        tasks = self.app.state.sections_service.tasks_from_blueprint(blueprint)
        self.assertIn(history["section_id"], {task["section_id"] for task in tasks})
        self.assertNotEqual("excluded_from_section_generation", history["automatic_resolution"]["action"])

    def test_optional_papers_do_not_bypass_outline_structure_or_matrix_validation(self) -> None:
        with TestClient(self.app) as client:
            selected = self.choose_outline(client, "custom")
            for outline in (
                "# No sections",
                "## 1. <!-- outline-untitled -->",
                "## History\nAssigned papers: P999.",
                "## History\nContext papers: P999.",
            ):
                with self.subTest(outline=outline):
                    response = client.put(
                        f"/api/v1/projects/{self.project_id}/planning/outline",
                        json={"revision": selected["matrix_revision"], "outline_style": "custom", "outline_md": outline},
                        headers=self.headers(),
                    )
                    self.assertEqual(422, response.status_code, response.text)
            saved = client.put(
                f"/api/v1/projects/{self.project_id}/planning/outline",
                json={"revision": selected["matrix_revision"], "outline_style": "custom", "outline_md": "## History\n\n## Mechanistic comparison\n"},
                headers=self.headers(),
            )
            self.assertEqual(200, saved.status_code, saved.text)
            self.assertTrue(saved.json()["outline_complete"])

    def test_reference_outline_is_registered(self) -> None:
        raw = "# Reference\n\n## 1. Mechanisms\nAssigned papers: P001.\nPurpose: compare.\n".encode()
        with patch.object(
            self.app.state.planning_service,
            "_analyze_reference_document",
            side_effect=self.isolated_reference_analysis,
        ):
            with TestClient(self.app) as client:
                current = self.planning(client)
                response = client.post(
                    f"/api/v1/projects/{self.project_id}/planning/reference-outlines",
                    json={
                        "revision": current["matrix_revision"],
                        "filename": "reference.md",
                        "content_base64": base64.b64encode(raw).decode(),
                    },
                    headers=self.headers(),
                )
                self.assertEqual(201, response.status_code, response.text)
                payload = self.planning(client)
                candidate_id = response.json()["candidate"]["candidate_id"]
                selected_response = client.put(
                    f"/api/v1/projects/{self.project_id}/planning/outline",
                    json={
                        "revision": payload["matrix_revision"],
                        "outline_style": f"reference:{candidate_id}",
                    },
                    headers=self.headers(),
                )
                self.assertEqual(200, selected_response.status_code, selected_response.text)
        candidate = response.json()["candidate"]
        self.assertTrue(candidate["source_artifact_id"])
        self.assertEqual("current_matrix_only", candidate["content_source"])
        self.assertFalse(candidate["reference_content_reused"])
        self.assertNotIn("Mechanisms", candidate["outline_md"])
        self.assertIn(candidate["candidate_id"], {item["candidate_id"] for item in payload["reference_outline_candidates"]})
        self.assertNotIn("Mechanisms", selected_response.json()["selected_outline_md"])

    def test_legacy_reference_candidate_fails_content_isolation(self) -> None:
        service = self.app.state.planning_service
        self.assertFalse(
            service._reference_candidate_is_isolated(
                {
                    "analysis_mode": "heading_extraction",
                    "outline_md": "## Source heading",
                }
            )
        )

    def test_reference_docx_content_is_not_used_as_candidate_headings(self) -> None:
        document = Document()
        document.add_heading("1. Mechanistic organization", level=1)
        document.add_paragraph("Reference discussion.")
        stream = BytesIO()
        document.save(stream)
        with patch.object(
            self.app.state.planning_service,
            "_analyze_reference_document",
            side_effect=self.isolated_reference_analysis,
        ):
            with TestClient(self.app) as client:
                current = self.planning(client)
                response = client.post(
                    f"/api/v1/projects/{self.project_id}/planning/reference-outlines",
                    json={
                        "revision": current["matrix_revision"],
                        "filename": "reference.docx",
                        "content_base64": base64.b64encode(stream.getvalue()).decode(),
                    },
                    headers=self.headers(),
                )
        self.assertEqual(201, response.status_code, response.text)
        candidate = response.json()["candidate"]
        self.assertEqual("ai_style_only_transfer_v2", candidate["analysis_mode"])
        self.assertNotIn("Mechanistic organization", candidate["outline_md"])
        self.assertIn("Copper allenation evidence", candidate["outline_md"])

    def test_blueprint_uses_current_matrix_and_outline(self) -> None:
        seed_verified_matrix(self.app.state.planning_service, self.first, self.project_id)
        with TestClient(self.app) as client:
            selected = self.choose_outline(client, "substrate")
            response = self.generate_candidate(client,
                json={"revision": 0},
                headers=self.headers(),
            )
            self.assertEqual(200, response.status_code, response.text)
            blueprint = response.json()["section_blueprint"]
        matrix_ids = {f"P{index:03d}" for index in range(1, 36)}
        assigned = {paper_id for section in blueprint["sections"] for paper_id in section["major_papers"]}
        self.assertTrue(assigned)
        self.assertLessEqual(assigned, matrix_ids)
        primary_occurrences = [
            paper_id
            for section in blueprint["sections"]
            for paper_id in section["primary_papers"]
        ]
        self.assertEqual(len(primary_occurrences), len(set(primary_occurrences)))
        introduction = next(
            section
            for section in blueprint["sections"]
            if section["section_role"] == "introduction"
        )
        self.assertFalse(any(section["section_role"] == "conclusion" for section in blueprint["sections"]))
        self.assertEqual([], introduction["major_papers"])
        self.assertTrue(introduction["supporting_papers"])
        self.assertTrue(
            all(section["writing_requirements"] for section in blueprint["sections"])
        )
        self.assertEqual(2, blueprint["schema_version"])
        body_sections = [
            section
            for section in blueprint["sections"]
            if section["section_role"] == "body" and section["primary_papers"]
        ]
        self.assertTrue(body_sections)
        # Abstract-only fixtures can guide a provisional plan without becoming
        # executable detailed claims or requiring an independent planning audit.
        self.assertTrue(all(section["thesis_status"] == "provisional" for section in body_sections))
        self.assertTrue(all(section["questions_to_answer"] and section["retrieval_directions"] for section in body_sections))
        self.assertTrue(all(not section["scientific_claims"] for section in body_sections))
        self.assertTrue(
            all(
                claim["claim_id"].startswith(f"{section['section_id']}-")
                and claim["primary_papers"]
                and claim["required_fact_roles"]
                for section in body_sections
                for claim in section["scientific_claims"]
            )
        )
        self.assertEqual([], introduction["scientific_claims"])
        self.assertTrue(
            all(
                section["review_claims"][0]["legacy_role"]
                == "writing_requirement"
                for section in blueprint["sections"]
            )
        )
        self.assertTrue(
            blueprint["paper_assignment_policy"][
                "introduction_and_conclusion_are_synthesis_only"
            ]
        )
        self.assertEqual(selected["outline_artifact_id"], blueprint["candidate_base_outline_artifact_id"])
        self.assertEqual(
            blueprint["classification_contract"]["fingerprint"],
            blueprint["classification_basis"]["axis_contract_fingerprint"],
        )
        self.assertEqual(
            blueprint["classification_contract"]["fingerprint"],
            blueprint["classification_contract_lineage"]["effective_fingerprint"],
        )
        structure_contract = blueprint["overview_structure_contract"]
        self.assertEqual("target_product", structure_contract["role"])
        self.assertEqual("resolved", structure_contract["status"])
        self.assertEqual("allene", structure_contract["motif"])
        self.assertEqual("*C=C=C*", structure_contract["smiles"])
        self.assertEqual("allenation", blueprint["rule_pack"])
        self.assertEqual(64, len(blueprint["rule_pack_sha256"]))

    def test_planning_bundle_exposes_scope_and_synthesis_requirements(self) -> None:
        with TestClient(self.app) as client:
            self.choose_outline(client, "reaction")
            generated = self.generate_candidate(client,
                json={"revision": 0},
                headers=self.headers(),
            )
        self.assertEqual(200, generated.status_code, generated.text)
        blueprint = generated.json()["section_blueprint"]
        self.assertTrue(blueprint["scope_diagnostics"]["can_confirm"])
        self.assertTrue(blueprint["taxonomy_diagnostics"]["can_confirm"])
        self.assertEqual(
            "reaction_strategy",
            blueprint["scope_contract"]["primary_navigation_axis"],
        )
        self.assertTrue(blueprint["synthesis_requirements"])
        self.assertTrue(
            all("academic_contract" in section for section in blueprint["sections"])
        )

    def test_catch_all_taxonomy_is_advisory_for_chapter_planning(self) -> None:
        seed_verified_matrix(self.app.state.planning_service, self.first, self.project_id)
        all_papers = ", ".join(f"P{index:03d}" for index in range(1, 36))
        outline = (
            "# Review\n\n"
            "## Introduction\nSection role: introduction\nPurpose: define scope.\n\n"
            "## Other or unspecified\nSection role: body\n"
            f"Assigned papers: {all_papers}.\n\n"
            "## Conclusion\nSection role: conclusion\nPurpose: synthesize.\n"
        )
        with TestClient(self.app) as client:
            selected = self.choose_outline(client, "custom")
            saved = client.put(
                f"/api/v1/projects/{self.project_id}/planning/outline",
                json={
                    "revision": selected["matrix_revision"],
                    "outline_style": "custom",
                    "outline_md": outline,
                },
                headers=self.headers(),
            )
            self.assertEqual(200, saved.status_code, saved.text)
            generated = self.generate_candidate(client,
                json={"revision": 0},
                headers=self.headers(),
            )
            self.assertEqual(200, generated.status_code, generated.text)
            blueprint = generated.json()["section_blueprint"]
            confirmed = client.post(
                f"/api/v1/projects/{self.project_id}/planning/blueprint/confirm",
                json={"revision": generated.json()["blueprint_revision"]},
                headers=self.headers(),
            )
        self.assertFalse(blueprint["taxonomy_diagnostics"]["can_confirm"])
        self.assertEqual(200, confirmed.status_code, confirmed.text)

    def test_scope_can_be_edited_with_the_existing_outline_save(self) -> None:
        with TestClient(self.app) as client:
            selected = self.choose_outline(client, "reaction")
            planning = self.planning(client)
            scope = dict(planning["scope_contract"])
            scope["target_question"] = "Which stereocontrol strategies are transferable?"
            saved = client.put(
                f"/api/v1/projects/{self.project_id}/planning/outline",
                json={
                    "revision": selected["matrix_revision"],
                    "outline_style": "reaction",
                    "outline_md": planning["selected_outline_md"],
                    "scope_contract": scope,
                },
                headers=self.headers(),
            )
            self.assertEqual(200, saved.status_code, saved.text)
            reloaded = self.planning(client)
        self.assertEqual(
            "Which stereocontrol strategies are transferable?",
            reloaded["scope_contract"]["target_question"],
        )
        self.assertEqual("user_edited", reloaded["scope_contract"]["source"])

    def test_duplicate_body_assignment_becomes_supporting_cross_reference(self) -> None:
        with TestClient(self.app) as client:
            selected = self.choose_outline(client, "custom")
            outline = (
                "# Review\n\n"
                "## Introduction\n"
                "Section role: introduction\n"
                "Purpose: define scope.\n\n"
                "## 1. First evidence theme\n"
                "Section role: body\n"
                "Assigned papers: P001, P002.\n\n"
                "## 2. Cross-cutting theme\n"
                "Section role: body\n"
                "Assigned papers: P001, P003.\n\n"
                "## Conclusion\n"
                "Section role: conclusion\n"
                "Purpose: synthesize findings.\n"
            )
            saved = client.put(
                f"/api/v1/projects/{self.project_id}/planning/outline",
                json={
                    "revision": selected["matrix_revision"],
                    "outline_style": "custom",
                    "outline_md": outline,
                },
                headers=self.headers(),
            )
            self.assertEqual(200, saved.status_code, saved.text)
            generated = self.generate_candidate(client,
                json={"revision": 0},
                headers=self.headers(),
            )
            self.assertEqual(200, generated.status_code, generated.text)
        sections = generated.json()["section_blueprint"]["sections"]
        first = next(section for section in sections if section["title"] == "First evidence theme")
        second = next(section for section in sections if section["title"] == "Cross-cutting theme")
        self.assertEqual(["P001", "P002"], first["primary_papers"])
        self.assertEqual(["P003"], second["primary_papers"])
        self.assertEqual(["P001"], second["supporting_papers"])

    def test_blueprint_confirmation_advances_to_sections(self):
        service = self.app.state.planning_service
        with TestClient(self.app) as client:
            seed_verified_matrix(service, self.first, self.project_id)
            self.choose_outline(client, "reaction")
            generated = self.generate_candidate(client, json={"revision": 0}, headers=self.headers()).json()
            self.assertEqual("completed", generated["section_blueprint"]["academic_planning"]["status"])
            self.assertIsNone(self.planning(client)["active_blueprint_artifact_id"])
            response = client.post(f"/api/v1/projects/{self.project_id}/planning/blueprint/confirm",
                json={"revision": 0, "artifact_id": generated["blueprint_artifact_id"]}, headers=self.headers())
            self.assertEqual(200, response.status_code, response.text)
            self.assertEqual("sections", client.get(f"/api/v1/projects/{self.project_id}").json()["current_stage"])

    def test_both_blueprint_routes_queue_the_same_idempotent_default_job(self):
        with patch.object(self.app.state.job_service, "execution_enabled", False), TestClient(self.app) as client:
            self.choose_outline(client, "reaction")
            headers = {**self.headers(), "Idempotency-Key": "default-argument-planning"}
            responses = [client.post(f"/api/v1/projects/{self.project_id}/planning/{route}",
                         json={"revision": 0}, headers=headers) for route in ("blueprint", "blueprint/jobs")]
            self.assertTrue(all(r.status_code == 202 for r in responses), [r.text for r in responses])
            self.assertEqual(responses[0].json()["id"], responses[1].json()["id"])
            self.assertIsNone(self.planning(client)["active_blueprint_artifact_id"])
            cancelled = client.post(f"/api/v1/jobs/{responses[0].json()['id']}/cancel", headers=self.headers())
            self.assertEqual(200, cancelled.status_code, cancelled.text)
            self.assertIsNone(self.planning(client)["active_blueprint_artifact_id"])

    def test_candidate_fact_outline_and_blueprint_commit_or_roll_back_together(self):
        from copy import deepcopy
        from review_writer_core.scientific_facts import review_fingerprint
        from review_writer_core.workflow.artifacts import MATRIX, PLANNING_OUTLINE, BLUEPRINT
        service, repo = self.app.state.planning_service, self.app.state.workflow_repository
        seed_verified_matrix(service, self.first, self.project_id)
        with TestClient(self.app) as client:
            self.choose_outline(client, "reaction")
            prepared = service.blueprint_job_payload(self.first, self.project_id, revision=0)
            self.assertTrue(prepared["integrated_fact_enrichment"]["enabled"])
            self.assertEqual(
                prepared["section_blueprint"]["source_matrix_artifact_id"],
                prepared["integrated_fact_enrichment"]["source_matrix_artifact_id"],
            )
            row = prepared["matrix_snapshot"]["rows"][0]
            fact = deepcopy(row["scientific_facts"][0])
            fact.update(fact_id="F-SUPPLEMENT", field_id="scope")
            fact["verification"]["input_fingerprint"] = review_fingerprint(fact)
            row["scientific_facts"].append(fact)
            built = offline_argument_planner(None, prepared)
            before = {name: (record.id if (record := repo.get_current_artifact(self.first.user_id, self.project_id, name)) else None)
                      for name in (MATRIX, PLANNING_OUTLINE, BLUEPRINT)}
            matrix_state = repo.get_stage_state(self.first.user_id, self.project_id, "matrix")
            candidate = service.publish_blueprint_candidate(self.first, self.project_id, built)
            first_candidate = candidate
            built["section_blueprint"]["recovery_note"] = "Revised argument, unchanged supplemented inputs"
            candidate = service.publish_blueprint_candidate(self.first, self.project_id, built)
            self.assertNotEqual(first_candidate["blueprint_artifact_id"], candidate["blueprint_artifact_id"])
            for key in ("source_matrix_artifact_id", "source_outline_artifact_id"):
                self.assertEqual(first_candidate["section_blueprint"][key], candidate["section_blueprint"][key])
            artifact = service.artifacts.resolve_owned_artifact(self.first.user_id, candidate["blueprint_artifact_id"]).artifact
            self.assertEqual({MATRIX: candidate["section_blueprint"]["source_matrix_artifact_id"],
                              PLANNING_OUTLINE: candidate["section_blueprint"]["source_outline_artifact_id"]},
                             artifact.metadata["planning_candidate_inputs"])
            matrix_artifact = service.artifacts.resolve_owned_artifact(self.first.user_id, candidate["section_blueprint"]["source_matrix_artifact_id"]).artifact
            self.assertNotEqual(artifact.producer_run_id, matrix_artifact.producer_run_id)
            from dataclasses import replace
            original_read = service._owned_blueprint_input
            def mismatched_registry(*args):
                document, record = original_read(*args)
                if record.id == artifact.id:
                    record = replace(record, metadata={"planning_candidate_inputs": {
                        MATRIX: before[MATRIX], PLANNING_OUTLINE: document["source_outline_artifact_id"]}})
                return document, record
            with patch.object(service, "_owned_blueprint_input", side_effect=mismatched_registry):
                with self.assertRaisesRegex(WorkflowValidationError, "registered planning artifact"):
                    service.confirm_blueprint(self.first, self.project_id, revision=0, artifact_id=candidate["blueprint_artifact_id"])
            snapshot = service.get(self.first, self.project_id)
            self.assertTrue(snapshot["blueprint_current"])
            self.assertIsNone(snapshot["active_blueprint_artifact_id"])
            self.assertIn("F-SUPPLEMENT", str(snapshot["blueprint_candidate_inputs"]))
            self.assertNotIn("F-SUPPLEMENT", str(snapshot["literature_matrix"]))
            original_upsert, writes = repo._upsert_current_artifact, []
            def fail_mid_transaction(*args, **kwargs):
                original_upsert(*args, **kwargs)
                writes.append(True)
                if len(writes) == 2:
                    raise RuntimeError("simulated interrupted promotion")
            with patch.object(repo, "_upsert_current_artifact", side_effect=fail_mid_transaction):
                with self.assertRaisesRegex(RuntimeError, "interrupted promotion"):
                    service.confirm_blueprint(self.first, self.project_id, revision=0, artifact_id=candidate["blueprint_artifact_id"])
            self.assertEqual(2, len(writes))
            after = {name: (record.id if (record := repo.get_current_artifact(self.first.user_id, self.project_id, name)) else None)
                     for name in before}
            self.assertEqual(before, after)
            self.assertEqual(matrix_state.revision, repo.get_stage_state(self.first.user_id, self.project_id, "matrix").revision)
            accepted = service.confirm_blueprint(self.first, self.project_id, revision=0, artifact_id=candidate["blueprint_artifact_id"])
            service.confirm_blueprint(self.first, self.project_id, revision=accepted["revision"], artifact_id=candidate["blueprint_artifact_id"])
            self.assertEqual(matrix_state.revision + 1, repo.get_stage_state(self.first.user_id, self.project_id, "matrix").revision)

    def test_stale_candidate_cannot_replace_newer_matrix_or_outline(self):
        service = self.app.state.planning_service
        seed_verified_matrix(service, self.first, self.project_id)
        with TestClient(self.app) as client:
            self.choose_outline(client, "reaction")
            generated = self.generate_candidate(client, json={"revision": 0}, headers=self.headers()).json()
            self.choose_outline(client, "substrate")
            with self.assertRaises(WorkflowConflict):
                service.confirm_blueprint(self.first, self.project_id, revision=0, artifact_id=generated["blueprint_artifact_id"])
            self.assertIsNone(service.get(self.first, self.project_id)["active_blueprint_artifact_id"])

    def test_unused_paper_remains_excluded_in_generation_display_and_confirmation(self):
        service = self.app.state.planning_service
        seed_verified_matrix(service, self.first, self.project_id)
        with TestClient(self.app) as client:
            self.choose_outline(client, "reaction")
            prepared = service.prepare_blueprint(self.first, self.project_id, revision=0)
            for section in prepared["section_blueprint"]["sections"]:
                for key in ("primary_papers", "major_papers", "supporting_papers", "context_papers"):
                    section[key] = [pid for pid in section.get(key) or [] if pid != "P002"]
            prepared["section_blueprint"]["sections"] = [s for s in prepared["section_blueprint"]["sections"]
                if s["section_role"] != "body" or s.get("primary_papers")]
            built = offline_argument_planner(None, prepared)
            self.assertIn("P002", built["section_blueprint"]["taxonomy_diagnostics"]["excluded_paper_ids"])
            # A legacy stored diagnostic must not contradict current documented exclusions.
            built["section_blueprint"]["taxonomy_diagnostics"]["issues"] = [{"rule_id": "taxonomy.orphan_papers",
                "severity": "planning_blocker", "paper_ids": ["P002"]}]
            candidate = service.publish_blueprint_candidate(self.first, self.project_id, built)
            page = self.planning(client)
            self.assertEqual([], page["taxonomy_diagnostics"]["orphan_paper_ids"])
            self.assertIn("P002", page["taxonomy_diagnostics"]["excluded_paper_ids"])
            self.assertIn("P002", [row["paper_id"] for row in page["literature_matrix"]["rows"]])
            confirmed = service.confirm_blueprint(self.first, self.project_id, revision=0, artifact_id=candidate["blueprint_artifact_id"])
            self.assertEqual("approved", confirmed["status"])

    def test_previous_blueprint_can_be_restored_as_a_new_review_version(self) -> None:
        seed_verified_matrix(self.app.state.planning_service, self.first, self.project_id)
        with TestClient(self.app) as client:
            self.choose_outline(client, "reaction")
            first = self.generate_candidate(client,
                json={"revision": 0},
                headers=self.headers(),
            )
            self.assertEqual(200, first.status_code, first.text)
            first_payload = first.json()
            approval = client.post(f"/api/v1/projects/{self.project_id}/planning/blueprint/confirm",
                json={"revision": 0, "artifact_id": first_payload["blueprint_artifact_id"]}, headers=self.headers())
            self.assertEqual(200, approval.status_code, approval.text)
            second = self.generate_candidate(client,
                json={"revision": approval.json()["revision"]},
                headers=self.headers(),
            )
            self.assertEqual(200, second.status_code, second.text)
            self.assertTrue(second.json()["candidate_pending"])
            snapshot = self.planning(client)
            self.assertEqual(first_payload["blueprint_artifact_id"], snapshot["active_blueprint_artifact_id"])
            self.assertEqual(approval.json()["revision"], second.json()["blueprint_revision"])
            accepted = client.post(
                f"/api/v1/projects/{self.project_id}/planning/blueprint/confirm",
                json={"revision": second.json()["blueprint_revision"], "artifact_id": second.json()["blueprint_artifact_id"]},
                headers=self.headers(),
            )
            self.assertEqual(200, accepted.status_code, accepted.text)
            restored = client.post(
                f"/api/v1/projects/{self.project_id}/planning/blueprint/restore",
                json={
                    "revision": accepted.json()["revision"],
                    "artifact_id": first_payload["blueprint_artifact_id"],
                },
                headers=self.headers(),
            )
            planning = self.planning(client)

        self.assertEqual(200, restored.status_code, restored.text)
        restored_payload = restored.json()
        self.assertEqual("review", restored_payload["status"])
        self.assertNotEqual(
            first_payload["blueprint_artifact_id"],
            restored_payload["blueprint_artifact_id"],
        )
        self.assertEqual(
            first_payload["blueprint_artifact_id"],
            planning["section_blueprint"]["restructure_record"][
                "restored_from_artifact_id"
            ],
        )

    def test_planning_contract_exposes_composite_tabs(self) -> None:
        with TestClient(self.app) as client:
            payload = self.planning(client)
        self.assertEqual(["matrix", "blueprint"], [tab["id"] for tab in payload["workspace"]["tabs"]])
        self.assertEqual("文献矩阵", payload["workspace"]["tabs"][0]["labels"]["zh"])
        self.assertEqual("Blueprint", payload["workspace"]["tabs"][1]["labels"]["en"])

    def test_planning_api_and_container_are_user_isolated(self) -> None:
        self.assertIs(
            self.app.state.planning_service,
            self.app.state.container.planning_service,
        )
        self.current = self.second
        with TestClient(self.app) as client:
            response = client.get(f"/api/v1/projects/{self.project_id}/planning")
        self.assertEqual(404, response.status_code, response.text)


if __name__ == "__main__":
    unittest.main()
